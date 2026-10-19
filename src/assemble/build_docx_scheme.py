"""Build outputs/mock_scheme_01.docx/.pdf through the same Word pipeline.

The scheme starts from the finished mock paper docx, so it keeps the same
template cover, embedded crest, header and footer. The question body is
removed and replaced by:
  - computed answers with step-by-step working (Caesar, Playfair, ACM,
    columnar; the other computed answers use the stored working),
  - the ACM matrix/ACL/CL and the rendered access-control diagram,
  - confidence labels and suggested mark allocation,
  - a table of the remaining theory placeholders.

Run:  uv run python -m src.assemble.build_docx_scheme
"""
from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.assemble.build_docx_paper import (
    DOT_BIN,
    OUT_DOCX as PAPER_DOCX,
    _acm_b_text,
    _acm_scenario_text,
    _qtext,
    build_slots,
    export_pdf,
    para_text,
    set_para_text,
)
from src.assemble.scheme_working import (
    confidence_for,
    mark_guidance_for,
    working_for,
)
from src.config import OUTPUTS_DIR
from src.generators.acm import (
    AccessScenario,
    build_matrix,
    derive_acls,
    derive_capabilities,
)

SCHEME_DOCX = OUTPUTS_DIR / "mock_scheme_01.docx"
SCHEME_PDF = OUTPUTS_DIR / "mock_scheme_01.pdf"

PLACEHOLDERS = [
    ("A Q1(a)", "attacks/malware", 2, "short objective"),
    ("A Q1(b)", "attacks/malware", 6, "define/explain"),
    ("A Q3(a)", "access control models", 2, "short objective"),
    ("A Q5(a)", "authentication/biometrics", 2, "short objective"),
    ("A Q6(a)", "software/OS security", 2, "short objective"),
    ("A Q6(b)", "software/OS security", 6, "define/explain"),
    ("A Q7(a)", "network security", 2, "short objective"),
    ("A Q7(b)", "network security", 6, "define/explain"),
    ("A Q8(a)", "trusted computing", 2, "short objective"),
    ("A Q8(b)", "trusted computing", 6, "define/explain"),
    ("A Q9(a)", "authentication/biometrics", 2, "short objective"),
    ("A Q9(b)", "authentication/biometrics", 6, "define/explain"),
    ("A Q10(a)", "attacks/malware", 2, "short objective"),
    ("A Q10(b)", "attacks/malware", 6, "define/explain"),
    ("B Q1(a)", "attacks/malware", 5, "scenario/justify"),
    ("B Q1(b)", "attacks/malware", 5, "scenario/justify"),
    ("B Q2(a)", "authentication/biometrics", 5, "scenario/justify"),
    ("B Q2(b)", "authentication/biometrics", 5, "scenario/justify"),
]


def _style_run(run, bold=False, italic=False, size=10.5, font="Arial") -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)


def add_para(doc, text: str = "", bold=False, italic=False, size=10.5,
             font="Arial", align=None, before=0, after=4,
             indent: float | None = None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        _style_run(p.add_run(text), bold=bold, italic=italic, size=size,
                   font=font)
    return p


def _set_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _cell_text(cell, text: str, bold=False, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _style_run(p.add_run(text), bold=bold, size=size)


def add_matrix_table(doc, scenario: AccessScenario, matrix) -> None:
    users, objects = scenario.users, scenario.objects
    table = doc.add_table(rows=1 + len(users), cols=1 + len(objects))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_borders(table)
    _cell_text(table.rows[0].cells[0], "Subject \\ Object", bold=True)
    for j, obj in enumerate(objects):
        _cell_text(table.rows[0].cells[j + 1], obj, bold=True)
    for i, user in enumerate(users):
        _cell_text(table.rows[i + 1].cells[0], user, bold=True)
        for j, obj in enumerate(objects):
            rights = matrix.get((user, obj), frozenset())
            _cell_text(table.rows[i + 1].cells[j + 1],
                       ", ".join(sorted(rights)) if rights else "-")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_acl_cl(doc, scenario: AccessScenario, matrix) -> None:
    acls = derive_acls(matrix, scenario.objects)
    caps = derive_capabilities(matrix, scenario.users)
    add_para(doc, "ACLs (per object):", bold=True, size=10, after=1)
    for obj in scenario.objects:
        entry = "; ".join(f"{u}: {', '.join(r)}"
                          for u, r in sorted(acls.get(obj, {}).items()))
        add_para(doc, f"  {obj}: {entry or '-'}", size=9.5, after=0)
    add_para(doc, "Capability lists (per subject):", bold=True, size=10,
             before=4, after=1)
    for user in scenario.users:
        entry = "; ".join(f"{o}: {', '.join(r)}"
                          for o, r in sorted(caps.get(user, {}).items()))
        add_para(doc, f"  {user}: {entry or '-'}", size=9.5, after=0)


def render_dot(dot: str, out_dir: Path) -> str:
    png = out_dir / "acm_diagram.png"
    proc = subprocess.run([DOT_BIN, "-Tpng", "-o", str(png)], input=dot,
                          text=True, capture_output=True, timeout=60)
    return str(png) if proc.returncode == 0 and png.exists() else ""


def _add_working(doc, working: str) -> None:
    lines = working.splitlines()
    if lines and lines[0].startswith("Answer:"):
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines = lines[1:]
    if lines and lines[0].strip() == "Working:":
        lines = lines[1:]
    add_para(doc, "Working:", bold=True, size=10, before=4, after=2)
    for line in lines:
        add_para(doc, line, font="Courier New", size=8.5, after=0, indent=0.4)


def _computed_header(doc, label: str, marks: int, confidence: str) -> None:
    add_para(doc, f"{label} - {marks} marks - {confidence}",
             bold=True, size=11, before=10, after=2)


def _question_text(doc, text: str) -> None:
    add_para(doc, "Question:", bold=True, size=10, before=2, after=1)
    add_para(doc, text, italic=True, size=9.5, after=2, indent=0.4)


def _answer(doc, answer: str) -> None:
    add_para(doc, f"Answer: {answer}", bold=True, size=10, before=3, after=1)


def build_scheme() -> None:
    slots = build_slots()
    computed = json.loads(
        (OUTPUTS_DIR / "computed_questions.json").read_text(encoding="utf-8"))
    acm_q = json.loads(
        (OUTPUTS_DIR / "acm_questions.json").read_text(encoding="utf-8"))[0]
    tmp = Path(tempfile.mkdtemp())
    diagram = render_dot(acm_q["dot"], tmp)

    doc = docx.Document(str(PAPER_DOCX))
    paras = doc.paragraphs

    # ---- cover: same template, relabelled as a marking scheme -------------
    title = next((p for p in paras if "UNIVERSITI TEKNOLOGI MARA" in
                  para_text(p._p) and "FINAL" in para_text(p._p)), None)
    if title is not None:
        set_para_text(title._p,
                      "UNIVERSITI TEKNOLOGI MARA\nMARKING SCHEME")

    cover_notes = [
        ("INSTRUCTIONS TO CANDIDATES", "MARKING SCHEME NOTES"),
        ("This question paper consists",
         "This marking scheme accompanies the CSC662 mock paper (October 2026)."),
        ("PART B (2 Questions)",
         "Computed answers show working; labels are computed / convention "
         "unverified / placeholder."),
        ("Answer ALL questions in the Answer Booklet",
         "Theory questions are placeholders - the theory module has not been built."),
        ("Do  not  bring",
         "Part A = 10 x (2 + 6); Part B = 2 x (5 + 5); total = 100 marks."),
        ("Please check to make sure",
         "Computed: 6 slots / 28 marks. Placeholders: 18 slots / 72 marks."),
        ("ii)an Answer Booklet",
         "Questions and answers are for local study use only."),
        ("Answer ALL questions in English",
         "See the review-pack checklist for Playfair, columnar and Vigenere "
         "convention checks."),
        ("DO NOT TURN THIS PAGE", "FOR MARKERS' USE ONLY"),
        ("This examination paper consists",
         "This marking scheme consists of N printed pages"),
    ]
    for prefix, new in cover_notes:
        p = next((p for p in doc.paragraphs
                  if para_text(p._p).strip().startswith(prefix)), None)
        if p is not None:
            set_para_text(p._p, new)

    # ---- remove the question body, keep cover + section formatting --------
    part_a = next(i for i, p in enumerate(doc.paragraphs)
                  if para_text(p._p).strip() == "PART A")
    end = next(i for i, p in enumerate(doc.paragraphs)
               if "END OF QUESTION PAPER" in para_text(p._p))
    body = doc.element.body
    doomed = [p._p for p in doc.paragraphs[part_a:end + 1]]
    for el in doomed:
        body.remove(el)

    # ---- scheme body ------------------------------------------------------
    add_para(doc, "CSC662 - MARKING SCHEME (MOCK OCTOBER 2026)",
             bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Confidence labels: computed / convention unverified / "
                  "placeholder. No official marking scheme exists for the real "
                  "CSC662 papers; this scheme is generated for the mock.",
             italic=True, size=9, after=10)
    add_para(doc, "Part A: 10 x (2 + 6) = 80 marks. Part B: 2 x (5 + 5) = "
                  "20 marks. Total: 100 marks.", size=10, after=8)

    labels = {
        "Q2a": ("PART A  Q2(a)", 2),
        "Q2b": ("PART A  Q2(b)", 6),
        "Q4a": ("PART A  Q4(a)", 2),
        "Q4b": ("PART A  Q4(b)", 6),
        "Q5b": ("PART A  Q5(b)", 6),
    }

    def add_acm() -> None:
        q = slots["Q3b"]
        scenario = AccessScenario(**q["scenario"])
        matrix = build_matrix(scenario)
        _computed_header(doc, "PART A  Q3(b)", 6,
                         "computed (consistency machine-verified)")
        _question_text(doc, _acm_scenario_text(q) + "\n" + _acm_b_text(q))
        add_para(doc, "Answer:", bold=True, size=10, before=3, after=2)
        add_matrix_table(doc, scenario, matrix)
        add_acl_cl(doc, scenario, matrix)
        if diagram:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(diagram, width=Cm(10))
            add_para(doc, "Figure: access-control diagram (answer).",
                     italic=True, size=9,
                     align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        else:
            add_para(doc, "Diagram rendering failed.", italic=True, size=9)
        add_para(doc, "Suggested marks: 2 matrix, 2 ACL, 1 capability list, "
                      "1 diagram.", italic=True, size=9, before=4, after=2)
        add_para(doc, "Confidence: computed - the matrix, ACLs and capability "
                      "lists rebuild each other without loss.",
                 italic=True, size=9, after=6)

    for key in ("Q2a", "Q2b", "Q3b", "Q4a", "Q4b", "Q5b"):
        if key == "Q3b":
            if slots.get("Q3b"):
                add_acm()
            continue
        q = slots.get(key)
        if not q:
            continue
        label, marks = labels[key]
        _computed_header(doc, label, marks, confidence_for(q))
        _question_text(doc, _qtext(q))
        _answer(doc, q["answer"])
        _add_working(doc, working_for(q))
        add_para(doc, "Suggested marks: " + mark_guidance_for(q),
                 italic=True, size=9, before=4, after=2)
        add_para(doc, f"Confidence: {confidence_for(q)}.",
                 italic=True, size=9, after=6)

    # ---- placeholders -----------------------------------------------------
    add_para(doc, "PLACEHOLDER SLOTS (18 slots, 72 marks)",
             bold=True, size=12, before=14, after=3)
    add_para(doc, "Theory questions are placeholders. The theory module has "
                  "not been built, so no model answers are provided for these "
                  "slots.", italic=True, size=9.5, after=4)
    table = doc.add_table(rows=1, cols=5)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_borders(table)
    for j, head in enumerate(("Slot", "Topic", "Marks", "Task", "Status")):
        _cell_text(table.rows[0].cells[j], head, bold=True)
    for slot, topic, marks, task in PLACEHOLDERS:
        cells = table.add_row().cells
        _cell_text(cells[0], slot)
        _cell_text(cells[1], topic)
        _cell_text(cells[2], str(marks))
        _cell_text(cells[3], task)
        _cell_text(cells[4], "placeholder")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_para(doc, "End of marking scheme.", italic=True, size=9,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=8)

    doc.save(str(SCHEME_DOCX))

    # page count -> cover line -> re-export
    pages = export_pdf(SCHEME_DOCX, SCHEME_PDF)
    d2 = docx.Document(str(SCHEME_DOCX))
    for p in d2.paragraphs:
        if para_text(p._p).strip().startswith(
                "This marking scheme consists of"):
            set_para_text(p._p,
                          f"This marking scheme consists of {pages} printed pages")
            break
    d2.save(str(SCHEME_DOCX))
    pages2 = export_pdf(SCHEME_DOCX, SCHEME_PDF)

    # checks
    import zipfile
    with zipfile.ZipFile(str(SCHEME_DOCX)) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    print(f"scheme: computed slots=6, placeholders={len(PLACEHOLDERS)} "
          f"({sum(p[2] for p in PLACEHOLDERS)} marks), pages={pages2} "
          f"(first export {pages}, match={pages == pages2})")
    print(f"scheme media: {media}")
    print(f"wrote {SCHEME_DOCX} and {SCHEME_PDF}")


if __name__ == "__main__":
    build_scheme()
