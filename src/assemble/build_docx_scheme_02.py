"""Build outputs/mock_scheme_02 through the same Word pipeline as paper 02.

The scheme keeps the paper's template cover, crest, header and footer, removes
the question body, and prints:
  - computed answers with working and confidence labels;
  - every accepted theory answer as `sourced`, with each marking point's
    verbatim quote and file + slide citation;
  - explicit placeholder status for unresolved theory slots.
"""
from __future__ import annotations

import json
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.assemble.build_docx_paper import export_pdf, para_text, set_para_text
from src.assemble.build_docx_scheme import (
    _add_working,
    _answer,
    _cell_text,
    _computed_header,
    _question_text,
    _set_table_borders,
    add_para,
)
from src.assemble.scheme_working import (
    confidence_for,
    mark_guidance_for,
    working_for,
)
from src.config import OUTPUTS_DIR

PAPER_DOCX = OUTPUTS_DIR / "mock_paper_02.docx"
SCHEME_DOCX = OUTPUTS_DIR / "mock_scheme_02.docx"
SCHEME_PDF = OUTPUTS_DIR / "mock_scheme_02.pdf"
PLAN_JSON = OUTPUTS_DIR / "planner_revised_realistic.json"
THEORY_JSON = OUTPUTS_DIR / "theory_drafts_realistic.json"


def _load() -> tuple[dict, dict, dict]:
    plan = json.loads(PLAN_JSON.read_text(encoding="utf-8"))
    slots = {s["slot"]: s for s in plan["slots"]}
    payload = json.loads(THEORY_JSON.read_text(encoding="utf-8"))
    accepted = {q["slot"]: q for q in payload.get("accepted", [])}
    placeholders = {p["slot"]: p for p in payload.get("placeholders", [])}
    return slots, accepted, placeholders


def _label(slot: dict, accepted: dict, placeholders: dict) -> str:
    if slot["kind"] == "computed":
        return "computed"
    if slot["slot"] in accepted:
        return "sourced"
    return "placeholder"


def _computed_question(slot: dict) -> dict | None:
    from src.assemble.build_docx_paper_02 import _computed_question as cq
    return cq(slot)


def _status(doc, slot: dict, accepted: dict, placeholders: dict) -> None:
    sid = slot["slot"]
    if slot["kind"] == "computed":
        q = _computed_question(slot)
        if q is None:
            add_para(doc, "Status: computed placeholder - no question "
                          "generated for this type.", italic=True, size=9)
            return
        _computed_header(doc, sid, slot["marks"], confidence_for(q))
        from src.assemble.build_docx_paper import paper_question
        _question_text(doc, paper_question(q, include_marks=False).strip())
        _answer(doc, q["answer"])
        if q.get("cipher") == "vigenere":
            add_para(doc, "The 26 x 26 Vigenere tableau is printed in the "
                          "question paper.", italic=True, size=9, after=2)
        _add_working(doc, working_for(q))
        add_para(doc, "Suggested marks: " + mark_guidance_for(q),
                 italic=True, size=9, before=4, after=2)
        add_para(doc, f"Confidence: {confidence_for(q)}.",
                 italic=True, size=9, after=6)
        return

    q = accepted.get(sid)
    if q is None:
        p = placeholders.get(sid)
        add_para(doc, f"{sid} - {slot['marks']} marks - PLACEHOLDER",
                 bold=True, size=11, before=10, after=2)
        add_para(doc, f"Topic: {slot['topic']}; command: {slot['task']}",
                 size=9.5, after=1)
        add_para(doc, "Status: placeholder - no sourced question passed "
                      "verification.", italic=True, size=9, after=1)
        if p:
            add_para(doc, "Reason: " + p["reason"], size=8.5, after=6)
        return

    add_para(doc, f"{sid} - {slot['marks']} marks - SOURCED",
             bold=True, size=11, before=10, after=2)
    _question_text(doc, q["question"])
    add_para(doc, "Marking points:", bold=True, size=10, before=3, after=2)
    table = doc.add_table(rows=1, cols=5)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_borders(table)
    for j, head in enumerate(("No.", "Point (1 mark)", "Evidence quote",
                              "Source", "Chunk")):
        _cell_text(table.rows[0].cells[j], head, bold=True, size=8)
    for i, pt in enumerate(q["marking_points"], 1):
        cells = table.add_row().cells
        _cell_text(cells[0], str(i), size=8)
        _cell_text(cells[1], pt["point"], size=8)
        _cell_text(cells[2], f'"{pt["quote"]}"', size=7.5)
        _cell_text(cells[3], f"{pt['source_file']}, slide "
                             f"{pt['source_slide']}", size=7.5)
        _cell_text(cells[4], pt["chunk_id"], size=7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_para(doc, "Confidence: sourced - each point carries a verbatim quote "
                  "from the cited lecture slide.", italic=True, size=9,
             before=2, after=1)
    if q.get("low_style_confidence"):
        add_para(doc, f"Style confidence: LOW - exact examples n<3; widened "
                      f"via {q.get('style_mode', '')}.", italic=True, size=9,
                 after=1)
    add_para(doc, "Citations: " + "; ".join(q.get("citations", [])),
             italic=True, size=9, after=6)


def build() -> None:
    slots, accepted, placeholders = _load()
    doc = docx.Document(str(PAPER_DOCX))
    paras = doc.paragraphs

    title = next((p for p in paras if "UNIVERSITI TEKNOLOGI MARA" in
                  para_text(p._p) and "FINAL" in para_text(p._p)), None)
    if title is not None:
        set_para_text(title._p,
                      "UNIVERSITI TEKNOLOGI MARA\nMARKING SCHEME")

    n_sourced = sum(1 for s in slots.values() if s["slot"] in accepted)
    n_ph = sum(1 for s in slots.values()
               if s["kind"] == "theory" and s["slot"] not in accepted)
    n_comp = sum(1 for s in slots.values() if s["kind"] == "computed")
    sourced_marks = sum(s["marks"] for s in slots.values()
                        if s["kind"] == "theory" and s["slot"] in accepted)

    cover_notes = [
        ("INSTRUCTIONS TO CANDIDATES", "MARKING SCHEME NOTES"),
        ("This question paper consists",
         "This marking scheme accompanies the CSC662 mock paper (October 2026)."),
        ("PART B (2 Questions)",
         "Computed answers show working; theory answers are sourced with a "
         "verbatim lecture quote and slide citation."),
        ("Answer ALL questions in the Answer Booklet",
         "Sourced theory marking points quote the lecture decks; any "
         "unresolved slot would remain a visible placeholder."),
        ("Do  not  bring",
         "Part A = 10 x (2 + 6); Part B = 2 x (5 + 5); total = 100 marks."),
        ("Please check to make sure",
         f"Computed: {n_comp} slot / 6 marks. Sourced theory: {n_sourced} "
         f"slots / {sourced_marks} marks; placeholders: {n_ph}."),
        ("ii)an Answer Booklet",
         "Questions and answers are for local study use only."),
        ("Answer ALL questions in English",
         "Near-copy rejection used all 35 past papers; style widening is "
         "recorded per slot."),
        ("DO NOT TURN THIS PAGE", "FOR MARKERS' USE ONLY"),
        ("This examination paper consists",
         "This marking scheme consists of N printed pages"),
    ]
    for prefix, new in cover_notes:
        p = next((p for p in doc.paragraphs
                  if para_text(p._p).strip().startswith(prefix)), None)
        if p is not None:
            set_para_text(p._p, new)

    part_a = next(i for i, p in enumerate(doc.paragraphs)
                  if para_text(p._p).strip().startswith("PART A"))
    end = next(i for i, p in enumerate(doc.paragraphs)
               if "END OF QUESTION PAPER" in para_text(p._p))
    body = doc.element.body
    for el in [p._p for p in doc.paragraphs[part_a:end + 1]]:
        body.remove(el)

    add_para(doc, "CSC662 - MARKING SCHEME (MOCK OCTOBER 2026)",
             bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Labels: computed / sourced / placeholder. No official "
                  "marking scheme exists for the real CSC662 papers; this is "
                  "generated for the mock.", italic=True, size=9, after=10)
    add_para(doc, f"Part A: 10 x (2 + 6) = 80 marks. Part B: 2 x (5 + 5) = "
                  f"20 marks. Total: 100 marks. Slots: {n_comp} computed, "
                  f"{n_sourced} sourced, {n_ph} placeholder.",
             size=10, after=8)

    for i in range(1, 11):
        add_para(doc, f"PART A - QUESTION {i}", bold=True, size=12,
                 before=12, after=2)
        _status(doc, slots[f"A{i}a"], accepted, placeholders)
        _status(doc, slots[f"A{i}b"], accepted, placeholders)

    for i, (a_id, b_id) in enumerate([("B1a", "B1b"), ("B2a", "B2b")],
                                    start=1):
        add_para(doc, f"PART B - QUESTION {i}", bold=True, size=12,
                 before=12, after=2)
        _status(doc, slots[a_id], accepted, placeholders)
        _status(doc, slots[b_id], accepted, placeholders)

    add_para(doc, "End of marking scheme.", italic=True, size=9,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=8)
    doc.save(str(SCHEME_DOCX))

    pages = export_pdf(SCHEME_DOCX, SCHEME_PDF)
    d2 = docx.Document(str(SCHEME_DOCX))
    for p in d2.paragraphs:
        if para_text(p._p).strip().startswith(
                "This marking scheme consists of"):
            set_para_text(p._p,
                          f"This marking scheme consists of {pages} printed pages")
            break
    d2.save(str(SCHEME_DOCX))
    pages2 = export_pdf(SCHEME_DOCX, SCHEME_PDF)
    print(f"scheme 02: computed={n_comp}, sourced={n_sourced}, "
          f"placeholders={n_ph}; pages={pages2} "
          f"(first export {pages}, match={pages == pages2})")
    print(f"wrote {SCHEME_DOCX} and {SCHEME_PDF}")


if __name__ == "__main__":
    build()
