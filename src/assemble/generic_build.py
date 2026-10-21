"""Parameterised paper/scheme builders used by the local web app.

The original builders read module-level paths for one seed.  This wrapper
serialises Word COM access and points the same builders at a run directory.
"""
from __future__ import annotations

import threading
from pathlib import Path

import docx
from docx.oxml.ns import qn

import src.assemble.build_docx_paper_04 as paper_mod
import src.assemble.build_docx_scheme_04 as scheme_mod
from src.assemble.build_docx_paper import export_pdf, para_text
from src.assemble.scheme_qa import SchemeQAError

_LOCK = threading.RLock()

_PUBLIC_REPLACEMENTS = (
    ("UNIVERSITI TEKNOLOGI MARA", "GENERIC UNIVERSITY"),
    ("Universiti Teknologi MARA", "Generic University"),
    ("CS/MOCK OCT 2026/CSC662", "MOCK EXAM 2026"),
    ("CS/MOCK", "MOCK EXAM"),
    ("2026/CSC662", ""),
)


def _replace_paragraph(p) -> None:
    ts = list(p.iter(qn("w:t")))
    if not ts:
        return
    joined = "".join(t.text or "" for t in ts).replace("\xa0", " ")
    new = joined
    for old, replacement in _PUBLIC_REPLACEMENTS:
        new = new.replace(old, replacement)
    if new != joined:
        ts[0].text = new
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            t.text = ""


def _replace_all(doc) -> None:
    for p in doc.element.body.iter(qn("w:p")):
        _replace_paragraph(p)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer):
            try:
                for p in part._element.iter(qn("w:p")):
                    _replace_paragraph(p)
            except Exception:
                continue


def _strip_crest(doc) -> None:
    """Remove header/cover graphics but keep later figures (e.g. tableau)."""
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer):
            try:
                paras = part.paragraphs
            except Exception:
                continue
            for p in paras:
                for tag in ("w:drawing", "w:pict"):
                    for el in p._p.findall(".//" + qn(tag)):
                        el.getparent().remove(el)
    stop = False
    for p in doc.paragraphs:
        if para_text(p._p).strip().startswith("PART A"):
            stop = True
        if stop:
            continue
        for tag in ("w:drawing", "w:pict"):
            for el in p._p.findall(".//" + qn(tag)):
                el.getparent().remove(el)


def publicise(docx_path: Path) -> None:
    doc = docx.Document(str(docx_path))
    _replace_all(doc)
    _strip_crest(doc)
    doc.save(str(docx_path))


def _reexport_pdf(docx_path: Path, pdf_path: Path) -> bool:
    try:
        export_pdf(docx_path, pdf_path)
        return pdf_path.exists()
    except Exception:
        if pdf_path.exists():
            pdf_path.unlink()
        return False


def build_paper(plan_json: Path, theory_json: Path, out_docx: Path,
                out_pdf: Path, public: bool = False) -> dict:
    """Build a paper; a missing Word/PDF backend is non-fatal."""
    plan_json = Path(plan_json).resolve()
    theory_json = Path(theory_json).resolve()
    out_docx = Path(out_docx).resolve()
    out_pdf = Path(out_pdf).resolve()
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    with _LOCK:
        paper_mod.PLAN_JSON = plan_json
        paper_mod.THEORY_JSON = theory_json
        paper_mod.OUT_DOCX = out_docx
        paper_mod.OUT_PDF = out_pdf
        error = ""
        applied_public = False
        try:
            applied_public = paper_mod.build(public=public)
        except Exception as e:  # Word COM may be unavailable
            error = str(e)
            if not out_docx.exists():
                raise
        if public and not applied_public:
            publicise(out_docx)
            pdf_ok = _reexport_pdf(out_docx, out_pdf)
        else:
            pdf_ok = out_pdf.exists()
    return {
        "docx": str(out_docx) if out_docx.exists() else "",
        "pdf": str(out_pdf) if pdf_ok else "",
        "pdf_error": "" if pdf_ok else (error or "PDF export unavailable"),
    }


def build_scheme(paper_docx: Path, plan_json: Path, theory_json: Path,
                 out_docx: Path, out_pdf: Path,
                 public: bool = False) -> dict:
    paper_docx = Path(paper_docx).resolve()
    plan_json = Path(plan_json).resolve()
    theory_json = Path(theory_json).resolve()
    out_docx = Path(out_docx).resolve()
    out_pdf = Path(out_pdf).resolve()
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    with _LOCK:
        scheme_mod.PAPER_DOCX = paper_docx
        scheme_mod.PLAN_JSON = plan_json
        scheme_mod.THEORY_JSON = theory_json
        scheme_mod.SCHEME_DOCX = out_docx
        scheme_mod.SCHEME_PDF = out_pdf
        error = ""
        try:
            scheme_mod.build(public=public)
        except SchemeQAError:
            raise
        except Exception as e:
            error = str(e)
            if not out_docx.exists():
                raise
        if public:
            publicise(out_docx)
            pdf_ok = _reexport_pdf(out_docx, out_pdf)
        else:
            pdf_ok = out_pdf.exists()
    return {
        "docx": str(out_docx) if out_docx.exists() else "",
        "pdf": str(out_pdf) if pdf_ok else "",
        "pdf_error": "" if pdf_ok else (error or "PDF export unavailable"),
    }
