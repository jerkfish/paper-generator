"""Automated scheme table/margin QA for Word and LibreOffice exports."""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from src.assemble.table_fit import (body_right_margin_twips,
                                   body_usable_width_twips, grid_errors,
                                   pdf_right_margin_errors)


class SchemeQAError(RuntimeError):
    pass


def find_soffice() -> str | None:
    env = os.environ.get("SOFFICE")
    if env and Path(env).exists():
        return env
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for path in (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ):
        if path.exists():
            return str(path)
    return None


def libreoffice_pdf(docx_path: Path, out_dir: Path) -> Path | None:
    soffice = find_soffice()
    if not soffice:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (Path(docx_path).stem + ".pdf")
    if out.exists():
        out.unlink()
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf",
         "--outdir", str(out_dir), str(docx_path)],
        capture_output=True, timeout=300)
    return out if out.exists() else None


def check_docx(docx_path: Path) -> list[str]:
    doc = Document(str(docx_path))
    usable = body_usable_width_twips(doc)
    errors = grid_errors(doc, usable)
    widths = []
    for tbl in doc.tables:
        widths.append(sum(int(c.get(qn("w:w")) or 0)
                          for c in tbl._tbl.tblGrid.findall(
                              qn("w:gridCol"))))
    print(f"usable_width_twips={usable}; tables={len(doc.tables)}; "
          f"grid_sums={sorted(set(widths))}", flush=True)
    return errors


def run_checks(docx_path: Path, word_pdf: Path,
               require_libreoffice: bool = True) -> dict:
    docx_path = Path(docx_path).resolve()
    word_pdf = Path(word_pdf).resolve()
    errors = check_docx(docx_path)
    doc = Document(str(docx_path))
    right_margin = body_right_margin_twips(doc)
    errors += pdf_right_margin_errors(word_pdf, right_margin)
    lo_pdf = None
    with tempfile.TemporaryDirectory() as td:
        lo_pdf = libreoffice_pdf(docx_path, Path(td))
        if lo_pdf is None:
            msg = "LibreOffice PDF export unavailable"
            if require_libreoffice:
                errors.append(msg)
            else:
                print("WARNING:", msg, flush=True)
        else:
            errors += pdf_right_margin_errors(lo_pdf, right_margin)
    if errors:
        raise SchemeQAError("scheme QA failed:\n- " + "\n- ".join(errors))
    print("scheme QA passed: grid widths and Word/LibreOffice margins",
          flush=True)
    return {"usable_width_twips": body_usable_width_twips(doc),
            "grid_errors": [], "margin_errors": []}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default="outputs/mock_scheme_04.docx")
    ap.add_argument("--word-pdf", default="outputs/mock_scheme_04.pdf")
    ap.add_argument("--no-libreoffice", action="store_true")
    args = ap.parse_args()
    print(run_checks(Path(args.docx), Path(args.word_pdf),
                     require_libreoffice=not args.no_libreoffice))


if __name__ == "__main__":
    main()
