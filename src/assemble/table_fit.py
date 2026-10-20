"""Word table geometry: usable width, fixed layout, grid and margin checks."""
from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# No., Point, Evidence quote, Source, Verdict, Overlap.
# Chunk was folded into Source (second line) to remove redundancy.
COLUMN_PROPORTIONS = (0.06, 0.30, 0.27, 0.16, 0.12, 0.09)


def twips(value) -> int:
    """python-docx Length -> twips (1 twip = 635 EMU)."""
    return int(round(int(value) / 635))


def body_usable_width_twips(doc) -> int:
    """Page width minus left/right margins for the body (last) section."""
    section = doc.sections[-1]
    return twips(section.page_width - section.left_margin
                 - section.right_margin)


def body_right_margin_twips(doc) -> int:
    return twips(doc.sections[-1].right_margin)


def column_widths(usable: int,
                  proportions: tuple[float, ...] = COLUMN_PROPORTIONS
                  ) -> list[int]:
    widths = [int(round(p * usable)) for p in proportions]
    widths[-1] = usable - sum(widths[:-1])
    return widths


def _set_tbl_props(tbl, usable: int) -> None:
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(usable))
    tblPr.append(tblW)
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:type"), "dxa")
    tblInd.set(qn("w:w"), "0")
    tblPr.append(tblInd)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    # Tight cell side margins give the narrow No./Overlap columns room.
    for el in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(el)
    mar = OxmlElement("w:tblCellMar")
    for side, width in (("top", 0), ("left", 60), ("bottom", 0),
                        ("right", 60)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(width))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _set_grid(tbl, widths: list[int]) -> None:
    grid = tbl._tbl.tblGrid
    for col in list(grid.findall(qn("w:gridCol"))):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def _row_props(row, header: bool) -> None:
    trPr = row._tr.get_or_add_trPr()
    for tag in ("w:cantSplit", "w:tblHeader"):
        for el in trPr.findall(qn(tag)):
            trPr.remove(el)
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def fix_tables(doc) -> int:
    """Apply fixed geometry to every table; returns the usable width."""
    usable = body_usable_width_twips(doc)
    widths = column_widths(usable)
    for tbl in doc.tables:
        _set_tbl_props(tbl, usable)
        _set_grid(tbl, widths)
        for ri, row in enumerate(tbl.rows):
            _row_props(row, header=(ri == 0))
            for ci, cell in enumerate(row.cells):
                if ci >= len(widths):
                    continue
                tcPr = cell._tc.get_or_add_tcPr()
                for el in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(el)
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:type"), "dxa")
                tcW.set(qn("w:w"), str(widths[ci]))
                tcPr.append(tcW)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(8)
    return usable


def grid_errors(doc, usable: int | None = None) -> list[str]:
    usable = usable or body_usable_width_twips(doc)
    errors = []
    for ti, tbl in enumerate(doc.tables, 1):
        cols = [int(c.get(qn("w:w")) or 0)
                for c in tbl._tbl.tblGrid.findall(qn("w:gridCol"))]
        total = sum(cols)
        if total > usable:
            errors.append(
                f"table {ti}: gridCol sum {total} > usable {usable} twips")
        tblW = tbl._tbl.tblPr.find(qn("w:tblW"))
        if tblW is not None:
            w = int(tblW.get(qn("w:w")) or 0)
            if w > usable:
                errors.append(
                    f"table {ti}: tblW {w} > usable {usable} twips")
    return errors


def pdf_right_margin_errors(pdf_path: Path, right_margin_twips: int,
                            dpi: int = 200,
                            tolerance_px: int = 6) -> list[str]:
    """Scan rendered pixels right of the right margin.

    White is >=245; antialiasing within `tolerance_px` of the boundary is
    ignored. Any darker pixel further right is reported with its page and x
    coordinate.
    """
    margin_pt = right_margin_twips / 20.0
    zoom = dpi / 72.0
    errors = []
    doc = pymupdf.open(str(pdf_path))
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom),
                              colorspace=pymupdf.csGRAY)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width)
        limit = int(round((page.rect.width - margin_pt) * zoom)) + tolerance_px
        if limit + 1 >= pix.width:
            continue
        band = img[:, limit + 1:]
        dark = np.argwhere(band < 245)
        if dark.size:
            max_x = int(dark[:, 1].max()) + limit + 1
            errors.append(
                f"page {i}: ink beyond right margin at x={max_x} "
                f"(limit {limit}px, page width {pix.width}px)")
    return errors
