"""Build the blind human-evaluation workbook."""
from __future__ import annotations

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.config import OUTPUTS_DIR
from src.planner.fingerprint_data import load_all_rows

EVAL_DIR = OUTPUTS_DIR / "eval_final_deepseek"
OUT = OUTPUTS_DIR / "human_eval.xlsx"


def _generated() -> list[dict]:
    out = []
    for p in sorted(EVAL_DIR.glob("seed_*_theory.json")):
        data = json.loads(p.read_text())
        seed = int(p.stem.split("_")[1])
        for q in data.get("accepted", []):
            out.append({"text": q["question"], "source": "generated",
                        "origin": f"seed {seed} / {q['slot']}",
                        "model": data.get("model", "")})
    return out


def _real() -> list[dict]:
    bad = ("confidential", "o hak cipta", "http", "universiti teknologi mara")
    out = []
    for r in load_all_rows():
        text = str(r["Original question text"]).replace("\n", " ").strip()
        if len(text.split()) < 8:
            continue
        if any(b in text.lower() for b in bad):
            continue
        out.append({"text": text, "source": "real",
                    "origin": f"{r['Paper ID']} Q{r['Q no.']}"
                              f"{r['Sub-part'] or ''}", "model": ""})
    return out


def main() -> None:
    rng = random.Random(20261020)
    gen = rng.sample(_generated(), 20)
    real = rng.sample(_real(), 10)
    items = gen + real
    rng.shuffle(items)
    wb = Workbook()
    ws = wb.active
    ws.title = "ratings"
    headers = ["id", "text", "answerable from the slides (1-5)",
               "feels like the course exam (1-5)", "unclear wording (Y/N)",
               "comments"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F0")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, item in enumerate(items, 1):
        qid = f"Q{i:03d}"
        item["id"] = qid
        ws.append([qid, item["text"], None, None, None, None])
        ws.cell(row=i + 1, column=2).alignment = Alignment(
            wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 95
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 38
    ws.freeze_panes = "A2"
    dv15 = DataValidation(type="whole", operator="between", formula1="1",
                          formula2="5", allow_blank=True)
    dvyn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv15)
    ws.add_data_validation(dvyn)
    dv15.add(f"C2:D{len(items) + 1}")
    dvyn.add(f"E2:E{len(items) + 1}")

    key = wb.create_sheet("key")
    key.append(["id", "source", "origin", "model"])
    for c in key[1]:
        c.font = Font(bold=True)
    for item in items:
        key.append([item["id"], item["source"], item["origin"],
                    item["model"]])
    key.sheet_state = "hidden"
    key.column_dimensions["A"].width = 10
    key.column_dimensions["B"].width = 12
    key.column_dimensions["C"].width = 28
    key.column_dimensions["D"].width = 28

    out = OUT
    wb.save(out)
    print(f"wrote {out}: {len(gen)} generated, {len(real)} real, shuffled")


if __name__ == "__main__":
    main()
