"""Step 4 - Read the fingerprint.

Computes the Topic x Task marks matrix for two views:
  (a) clean papers only (computed sum in [98,102]), broken down by era,
  (b) all papers.
Plus: machine-verifiable marks per paper (mean + per era), repeated questions
across years (normalised text), excluded papers, whole-block rows, and the
"other" share. Writes outputs/fingerprint_report.md.

Run:  uv run python -m src.fingerprint.fingerprint
"""
from __future__ import annotations

import re
from collections import defaultdict

from openpyxl import load_workbook

from src.config import OUTPUTS_DIR, TRACKING_XLSX
from src.fingerprint.types import (TASK_NAMES, TOPIC_NAMES, is_ac_construction,
                                   is_whole_block_row, label_group)

CONSTRUCTION_TOPIC = "access control (construction)"
ALL_TOPIC_KEYS = list(TOPIC_NAMES)
ALL_TOPIC_KEYS[ALL_TOPIC_KEYS.index("access control models")] = "access control models"
ALL_TOPIC_KEYS.insert(ALL_TOPIC_KEYS.index("access control models") + 1, CONSTRUCTION_TOPIC)

CLEAN_RANGE = (98, 102)
MATRIX_TASKS = TASK_NAMES


def _sheet(name: str) -> list[dict]:
    ws = load_workbook(TRACKING_XLSX)[name]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def _era_of(pid: str, papers: dict[str, str]) -> str:
    return papers.get(pid, "?").split(" ")[0]


def _group_key(row: dict) -> tuple[str, str]:
    q = str(row["Q no."])
    m = re.match(r"^([A-Z])?(\d+)", q)
    root = (m.group(1) or "") + (m.group(2) or q) if m else q
    return (row["Paper ID"], root)


def label_all(rows: list[dict]) -> list[dict]:
    """Fresh Topic x Task labels for every row, with sibling inheritance."""
    order: dict[tuple[str, str], list[int]] = defaultdict(list)
    for i, row in enumerate(rows):
        order[_group_key(row)].append(i)
    cache: dict[tuple[str, str], list[dict]] = {}
    for key, idxs in order.items():
        cache[key] = label_group([rows[i] for i in idxs])
    for key, idxs in order.items():
        for i, lab in zip(idxs, cache[key]):
            rows[i]["_lab"] = lab
    return rows


def _topic_key(row: dict, lab: dict) -> str:
    if lab["topic"] == "access control models" and \
            is_ac_construction(str(row["Original question text"]), lab["topic"]):
        return CONSTRUCTION_TOPIC
    return lab["topic"]


def matrix_for(rows: list[dict]) -> dict[tuple[str, str], dict]:
    m: dict[tuple[str, str], dict] = defaultdict(lambda: {"marks": 0.0, "rows": 0,
                                                          "mv_marks": 0.0})
    for r in rows:
        if is_whole_block_row(r):
            continue
        lab = r["_lab"]
        key = (_topic_key(r, lab), lab["task"])
        cell = m[key]
        cell["rows"] += 1
        if r["Marks"] is not None:
            cell["marks"] += float(r["Marks"])
            if lab["machine_verifiable"]:
                cell["mv_marks"] += float(r["Marks"])
    return m


def _norm_text(text: str) -> str:
    t = str(text).lower()
    t = re.sub(r"\(\d+(?:\.\d+)?\s*marks?\)", "", t)
    t = re.sub(r"\d+", "#", t)
    t = re.sub(r"[^a-z#]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def repeated_questions(rows: list[dict]) -> list[dict]:
    by_paper: dict[str, dict[str, tuple]] = defaultdict(dict)
    for r in rows:
        if is_whole_block_row(r):
            continue
        norm = _norm_text(r["Original question text"])
        if len(norm) < 50:
            continue
        if norm not in by_paper[r["Paper ID"]]:
            by_paper[r["Paper ID"]][norm] = r
    seen: dict[str, dict] = {}
    for pid, norms in by_paper.items():
        for norm, row in norms.items():
            entry = seen.setdefault(norm, {"papers": set(), "rows": [],
                                           "snippet": str(row["Original question text"]),
                                           "marks": row["Marks"]})
            entry["papers"].add(pid)
            entry["rows"].append(row)
    out = [{"norm": norm, "count": len(e["papers"]),
            "years": sorted({int(re.search(r"-(\d{4})-", p).group(1))
                             for p in e["papers"]}),
            "snippet": e["snippet"], "marks": e["marks"]}
           for norm, e in seen.items() if len(e["papers"]) >= 2]
    out.sort(key=lambda e: (-e["count"], -len(e["years"])))
    return out


def fmt_matrix(m: dict[tuple[str, str], dict], title: str) -> str:
    total = sum(c["marks"] for c in m.values())
    lines = [f"\n## {title}  (total {total:.0f} marks)\n",
             "| Topic | " + " | ".join(MATRIX_TASKS) + " | topic total | share |"]
    lines.append("|---|" + "---|" * (len(MATRIX_TASKS) + 1))
    for topic in [t for t in ALL_TOPIC_KEYS if any(k[0] == t for k in m)]:
        row_total = sum(m[(topic, t)]["marks"] for t in MATRIX_TASKS)
        cells = []
        for t in MATRIX_TASKS:
            cell = m.get((topic, t))
            if not cell or cell["marks"] == 0:
                cells.append("·")
            elif cell["mv_marks"] > 0:
                cells.append(f"{cell['marks']:.0f} ({cell['mv_marks']:.0f}v)")
            else:
                cells.append(f"{cell['marks']:.0f}")
        share = 100 * row_total / total if total else 0
        lines.append(f"| {topic} | " + " | ".join(cells) +
                     f" | {row_total:.0f} | {share:.0f}% |")
    return "\n".join(lines)


def mv_per_paper(rows: list[dict], papers_era: dict[str, str]) -> str:
    per: dict[str, float] = defaultdict(float)
    era_rows: dict[str, list[float]] = defaultdict(list)
    for r in rows:
        if is_whole_block_row(r):
            continue
        if r["_lab"]["machine_verifiable"] and r["Marks"] is not None:
            per[r["Paper ID"]] += float(r["Marks"])
    for pid, marks in per.items():
        era_rows[_era_of(pid, papers_era)].append(marks)
    lines = ["\n## Machine-verifiable marks per paper\n"]
    lines.append("| Era | MV marks per paper (per-paper values) | mean |")
    lines.append("|---|---|---|")
    order = ["classic", "transitional", "modern"]
    all_vals = []
    for era in order:
        vals = sorted(era_rows.get(era, []), reverse=True)
        all_vals += vals
        mean = sum(vals) / len(vals) if vals else 0.0
        lines.append(f"| {era} (based on {len(vals)} papers with MV marks) | "
                     f"{', '.join(f'{v:.0f}' for v in vals)} | {mean:.1f} |")
    total_mean = sum(all_vals) / len(all_vals) if all_vals else 0.0
    lines.append(f"| **all** (based on {len(all_vals)} papers) | | **{total_mean:.1f}** |")
    lines.append(f"\nTotal machine-verifiable marks: {sum(all_vals):.0f}")
    return "\n".join(lines)


def main() -> None:
    rows = _sheet("Questions")
    checks = _sheet("Paper marks check")
    papers = {r["Paper ID"]: r["Era"] for r in _sheet("Papers")}
    label_all(rows)

    check_map = {c["Paper ID"]: c for c in checks}
    clean = {pid for pid, c in check_map.items()
             if isinstance(c["Computed sum"], (int, float))
             and CLEAN_RANGE[0] <= c["Computed sum"] <= CLEAN_RANGE[1]}
    excluded = [(pid, c["Computed sum"]) for pid, c in check_map.items() if pid not in clean]
    excluded.sort()

    clean_rows = [r for r in rows if r["Paper ID"] in clean]
    clean_by_era: dict[str, list[dict]] = defaultdict(list)
    for r in clean_rows:
        clean_by_era[_era_of(r["Paper ID"], papers)].append(r)

    wb_rows = [r for r in rows if is_whole_block_row(r)]
    wb_marks = sum(float(r["Marks"]) for r in wb_rows if r["Marks"] is not None)

    # "other" share
    total_marks = sum(float(r["Marks"]) for r in rows
                      if not is_whole_block_row(r) and r["Marks"] is not None)
    other_marks = sum(float(r["Marks"]) for r in rows
                      if not is_whole_block_row(r) and r["Marks"] is not None
                      and r["_lab"]["topic"] == "other")
    other_share = 100 * other_marks / total_marks if total_marks else 0.0

    report = ["# CSC662 fingerprint (Topic x Task, marks)\n",
              "Computed by code from the Questions sheet. Marks only; 'v' in a cell "
              "= machine-verifiable marks. Whole-block fallback rows are excluded "
              "from every table and listed separately."]

    for era in ("classic", "transitional", "modern"):
        rows_e = clean_by_era[era]
        n_papers = len({r["Paper ID"] for r in rows_e})
        report.append(fmt_matrix(matrix_for(rows_e),
                                 f"View (a) clean papers - {era} era "
                                 f"(based on {n_papers} papers)"))
    all_n = len({r["Paper ID"] for r in rows})
    report.append(fmt_matrix(matrix_for(rows), f"View (b) all papers (based on {all_n} papers)"))

    report.append(mv_per_paper(rows, papers))

    prop = [r for r in rows
            if r["_lab"].get("machine_verifiable")
            and any("inputs inherited" in f for f in r["_lab"].get("flags", []))]
    prop_marks = sum(float(r["Marks"]) for r in prop if r["Marks"] is not None)
    report.append(f"\n## RSA/DH parameter propagation (owner rule #5)\n"
                  f"{len(prop)} rows became machine-verifiable because p/q/e/key "
                  f"parameters were inherited from sibling rows "
                  f"({prop_marks:.0f} printed marks).")

    report.append("\n## Repeated questions across years (normalised text, "
                  f"based on {all_n} papers)")
    report.append("\n| #papers | years | marks | question |")
    report.append("|---|---|---|---|")
    for e in repeated_questions(rows)[:15]:
        years = ",".join(str(y) for y in e["years"])
        snip = re.sub(r"\s+", " ", e["snippet"])[:110]
        report.append(f"| {e['count']} | {years} | {e['marks']} | {snip} |")

    report.append(f"\n## Other-share check\n")
    report.append(f"'other' topic rows: {other_marks:.0f} / {total_marks:.0f} marks "
                  f"= {other_share:.1f}% (aim < 15%)")

    report.append(f"\n## Excluded from view (a) - computed sums outside [{CLEAN_RANGE[0]}.."
                  f"{CLEAN_RANGE[1]}]")
    for pid, total in excluded:
        report.append(f"- {pid}: sum = {total}")

    report.append(f"\n## Whole-block fallback rows (excluded from statistics)")
    report.append(f"{len(wb_rows)} rows, {wb_marks:.0f} marks (verbatim check failed; "
                  f"stored as whole question blocks)")

    OUTPUTS_DIR.mkdir(exist_ok=True)
    out_path = OUTPUTS_DIR / "fingerprint_report.md"
    out_path.write_text("\n".join(report) + "\n", encoding="utf-8")

    print(f"wrote {out_path}")
    for section in report:
        print(section)


if __name__ == "__main__":
    main()
