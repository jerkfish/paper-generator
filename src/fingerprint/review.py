"""Review aid for the Topic x Task relabel: 20 random rows per task for the
owner to check, plus the code-defined machine-verifiable summary per topic.

Run:  uv run python -m src.fingerprint.review
"""
from __future__ import annotations

import random
from collections import defaultdict

from openpyxl import load_workbook

from src.config import TRACKING_XLSX
from src.fingerprint.types import TASK_NAMES, TOPIC_NAMES, classify_row


def load_rows() -> list[dict]:
    ws = load_workbook(TRACKING_XLSX)["Questions"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def main() -> None:
    rows = load_rows()
    labelled = [(classify_row(str(r["Original question text"]), r["Q no."],
                              str(r["Question type"])), r) for r in rows]

    rng = random.Random(42)
    for task in TASK_NAMES:
        pool = [(lab, r) for lab, r in labelled if lab["task"] == task]
        print(f"\n===== TASK: {task} ({len(pool)} rows) =====")
        for lab, r in rng.sample(pool, min(20, len(pool))):
            snippet = str(r["Original question text"])[:95].replace("\n", " ")
            mv = "Y" if lab["machine_verifiable"] else "n"
            print(f"  [{lab['topic']:20} mv={mv}] {r['Paper ID'][6:]} "
                  f"Q{r['Q no.']}{r['Sub-part']} ({r['Marks']}m): {snippet}")

    print("\n===== MACHINE-VERIFIABLE SUMMARY BY TOPIC (code-defined) =====")
    stats = defaultdict(lambda: {"rows": 0, "marks": 0.0, "mv_rows": 0,
                                 "mv_marks": 0.0, "no_marks": 0})
    for lab, r in labelled:
        s = stats[lab["topic"]]
        s["rows"] += 1
        if r["Marks"] is None:
            s["no_marks"] += 1
        else:
            s["marks"] += float(r["Marks"])
        if lab["machine_verifiable"]:
            s["mv_rows"] += 1
            if r["Marks"] is not None:
                s["mv_marks"] += float(r["Marks"])
    print(f"{'topic':22} {'rows':>5} {'marks':>8} {'mv rows':>8} {'mv marks':>9} "
          f"{'mv%marks':>9}")
    for topic in TOPIC_NAMES:
        s = stats[topic]
        if s["rows"] == 0:
            continue
        pct = 100 * s["mv_marks"] / s["marks"] if s["marks"] else 0.0
        marks_s = f"{s['marks']:.0f}" + ("?" if s["no_marks"] else "")
        print(f"{topic:22} {s['rows']:5} {marks_s:>8} {s['mv_rows']:8} "
              f"{s['mv_marks']:9.0f} {pct:8.0f}%")


if __name__ == "__main__":
    main()
