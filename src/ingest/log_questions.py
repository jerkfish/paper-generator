"""Step 3 - Log questions.

For every paper (excluding byte-identical "(2)" copies): extract the text
(text layer, or full Windows OCR for scans - cached), split it into question
blocks with code, ask the LLM to propose per-question/sub-part labels while
quoting the ORIGINAL wording verbatim, verify the quotes against the source
text, and write the "Questions" sheet plus a per-paper "Paper marks check"
sheet (computed sum vs official total; mismatches are listed, not hidden).

Run:  uv run python -m src.ingest.log_questions [--only all|text|scanned]
"""
from __future__ import annotations

import argparse
import pathlib
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from src import config
from src.ingest import ocr as ocr_mod
from src.ingest.models import PaperQuestions
from src.llm import chat_json

QUESTION_LINE_PAT = re.compile(r"^\s*QUESTION\s+(\d+)\b", re.I)
PART_LINE_PAT = re.compile(r"^\s*(PART\s+([AB]))\b", re.I)
MARKS_PAT = re.compile(r"\((\d+(?:\.\d+)?)\s*marks?\)", re.I)

QUESTION_TYPES = [
    "Objective (MCQ)", "Short answer", "Explanation/Theory", "Calculation",
    "Diagram/Label", "Case study", "Comparison/Discussion", "Other",
]
OCR_MARKERS = [
    re.compile(r"UNIVERSITI", re.I), re.compile(r"\bCSC\s*662\b", re.I),
    re.compile(r"QUESTION", re.I), re.compile(r"MARK", re.I),
]

QUESTIONS_HEADERS = [
    "Paper ID", "Q no.", "Sub-part", "Topic", "Question type", "Marks",
    "Has diagram", "Diagram type", "Machine-verifiable", "Wording note",
    "Original question text", "Confidence",
]
CHECK_HEADERS = ["Paper ID", "Official total", "Computed sum", "Status", "Note"]

# openpyxl rejects these control chars (OCR text contains a few)
_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean(value):
    return _ILLEGAL_RE.sub("", value) if isinstance(value, str) else value

SYSTEM_PROMPT = f"""You label past-exam questions. You receive one exam paper split into question blocks.
For every question block, split it into sub-parts (a, b, c... or i, ii...) ONLY when each
sub-part carries its own marks; otherwise one row for the whole question.
Copy the question text VERBATIM from the block into original_text - do not paraphrase,
repair spelling, or translate. Keep every mark figure exactly as printed.
Output exactly one JSON object with this shape:
{{"questions": [{{"q_no": "1", "sub_part": "a", "original_text": "...verbatim...",
"marks": 10, "topic": "short topic e.g. RSA, access control, DES",
"question_type": "one of: {", ".join(QUESTION_TYPES)}",
"has_diagram": false, "diagram_type": "", "machine_verifiable": "Y",
"wording_note": "short note on phrasing style (e.g. 'state and justify')"}}]}}
Rules: q_no must be copied EXACTLY from the block label, including any part letter
(e.g. "A1" or "B2" - never drop or invent the part letter). sub_part is "" for a whole
question. marks is a number or null when no marks are printed for that row.
machine_verifiable is "Y" when a program could compute the exact answer (cipher steps,
numeric crypto, fixed-fact MCQ), "Partial" when only part is computable, "N" for
essays/opinions. has_diagram true only if the paper shows a figure/table/diagram inside
this question."""


def _norm(s: str) -> str:
    # PDF text layers often hold curly quotes/dashes - compare against straight
    s = (s.replace("“", '"').replace("”", '"')
         .replace("‘", "'").replace("’", "'")
         .replace("–", "-").replace("—", "-"))
    return re.sub(r"\s+", "", s)


def _qkey(q: str) -> tuple:
    """Canonical question key: ('A', 3) for 'A3', ('', 3) for '3', else raw."""
    m = re.match(r"^([AB])?(\d+)$", str(q).strip())
    if m:
        return (m.group(1) or "", int(m.group(2)))
    return (str(q).strip(), None)


def _fuzzy_contains(q: str, block: str, threshold: float = 0.9) -> bool:
    """Exact substring, or >=threshold of the quote matches contiguously."""
    from difflib import SequenceMatcher

    if q in block:
        return True
    match = SequenceMatcher(None, q, block).find_longest_match(0, len(q), 0, len(block))
    return match.size >= threshold * len(q)


def _full_text(path: pathlib.Path) -> str:
    from pypdf import PdfReader

    parts = []
    try:
        for page in PdfReader(str(path)).pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                pass
    except Exception:
        return ""
    return "\n".join(parts)


def _split_blocks(text: str) -> list[tuple[str, str]]:
    """Split paper text into (q_no, block) pairs, tracking Part A/B where present."""
    # OCR text often has no line breaks at all - force breaks before markers
    text = re.sub(r"(?<![A-Za-z0-9])QUESTION\s+\d+\b",
                  lambda m: "\n" + m.group(0), text, flags=re.I)
    text = re.sub(r"(?<![A-Za-z0-9])PART\s+[AB]\b",
                  lambda m: "\n" + m.group(0), text, flags=re.I)
    has_parts = bool(PART_LINE_PAT.search(text))
    blocks: list[tuple[str, str]] = []
    cur_q: str | None = None
    cur_part = ""
    buf: list[str] = []

    def flush() -> None:
        nonlocal buf
        if cur_q is not None:
            blocks.append((cur_q, "\n".join(buf)))
        buf = []  # always drop the buffer - cover preamble must not leak into Q1

    for line in text.splitlines():
        pm = PART_LINE_PAT.match(line)
        if pm:
            flush()
            cur_part = pm.group(2).upper()
            buf.append(line)
            continue
        qm = QUESTION_LINE_PAT.match(line)
        if qm:
            flush()
            cur_q = (cur_part if has_parts and cur_part else "") + qm.group(1)
        buf.append(line)
    flush()
    if not blocks:  # no "QUESTION n" markers found - hand the whole text to the LLM
        return [("?", text)]

    # fold short header-only tails ("PART B (20 MARKS)") into the previous block
    merged: list[tuple[str, str]] = []
    for q_no, body in blocks:
        if merged and not QUESTION_LINE_PAT.search(body) and len(_norm(body)) < 80:
            pq, pbody = merged[-1]
            merged[-1] = (pq, pbody + "\n" + body)
        else:
            merged.append((q_no, body))
    blocks = merged

    # numbering restarts => the paper has parts; prefix A, B, ... per run
    nums = [int(q) if re.fullmatch(r"\d+", q) else None for q, _ in blocks]
    has_resets = any(nums[i] is not None and nums[i - 1] is not None
                     and nums[i] <= nums[i - 1] for i in range(1, len(nums)))
    if not has_resets:
        return blocks
    out: list[tuple[str, str]] = []
    part, prev = "", None
    for q_no, body in blocks:
        if re.fullmatch(r"\d+", q_no):
            n = int(q_no)
            if prev is None:
                part = "A"
            elif n <= prev:
                part = chr(ord(part) + 1)
            prev = n
            out.append((part + q_no, body))
        else:
            out.append((q_no, body))
    return out


def _block_marks(block: str) -> float | None:
    values = [float(m) for m in MARKS_PAT.findall(block)]
    if len(values) >= 2 and abs(values[0] - sum(values[1:])) < 0.01:
        # header total ("(10 Marks)") + sub-part marks - don't double count
        values = values[1:]
    return sum(values) if values else None


def _label_paper(pid: str, blocks: list[tuple[str, str]]) -> list[dict]:
    """One LLM call per paper: label every block at once."""
    parts = [f"=== QUESTION BLOCK {q_no} ===\n{block[:2500]}" for q_no, block in blocks]
    user = f"Paper {pid}.\n\n" + "\n\n".join(parts)
    try:
        data = chat_json(SYSTEM_PROMPT, user)
        parsed = PaperQuestions(paper_id=pid, questions=data.get("questions", []))
        return [q.model_dump() for q in parsed.questions]
    except (ValidationError, RuntimeError, KeyError) as e:
        print(f"    !! LLM label failed for {pid}: {e}", flush=True)
        return []


def _verify_block(block: str, rows: list[dict], full_paper_norm: str = "",
                  q_no: str = "?") -> tuple[list[dict], bool]:
    """Check every quote against the paper text; fall back to a whole-block row
    only when a quote cannot be matched anywhere in the paper."""
    norm_block = _norm(block)
    quoted_chars = 0
    matched = 0
    for row in rows:
        q = _norm(str(row["original_text"]))
        if not q:
            continue
        where = ""
        if q in norm_block:
            where = "exact"
        elif _fuzzy_contains(q, norm_block):
            where = "fuzzy"
        elif full_paper_norm and (q in full_paper_norm
                                  or _fuzzy_contains(q, full_paper_norm, threshold=0.85)):
            where = "elsewhere"
        if where:
            matched += 1
            if where == "exact":
                quoted_chars += len(q)
            elif where == "fuzzy":
                quoted_chars += len(q)
                row["wording_note"] = (row.get("wording_note") + "; "
                                       if row.get("wording_note") else "") \
                                       + "quote fuzzy-matched to paper"
            else:
                row["wording_note"] = (row.get("wording_note") + "; "
                                       if row.get("wording_note") else "") \
                                       + "quote matched elsewhere in paper (block split LLM-proposed)"
    coverage = quoted_chars / max(len(norm_block), 1)
    if matched == len(rows) and rows:
        # fill missing marks from the "(N marks)" printed right after each quote
        for row in rows:
            if row.get("marks") is not None:
                continue
            q = _norm(str(row["original_text"]))
            pos = norm_block.find(q)
            if pos == -1:
                continue
            m = re.search(r"\((\d+(?:\.\d+)?)marks?\)",
                          norm_block[pos + len(q):pos + len(q) + 400])
            if m:
                row["marks"] = float(m.group(1))
                row["wording_note"] = (row.get("wording_note") + "; "
                                       if row.get("wording_note") else "") \
                                       + "marks computed by code from '(N marks)' after quote"
        if coverage < 0.5:
            rows[0]["wording_note"] = (rows[0].get("wording_note") + "; "
                                       if rows[0].get("wording_note") else "") \
                                       + "low block coverage - check sub-parts are complete"
        return rows, True
    # a quote could not be matched anywhere - keep the block as one row
    return [{
        "q_no": q_no,
        "sub_part": "",
        "original_text": block.strip(),
        "marks": _block_marks(block),
        "topic": "unsure",
        "question_type": "Other",
        "has_diagram": False,
        "diagram_type": "",
        "machine_verifiable": "N",
        "wording_note": "verbatim check failed; whole block stored",
    }], False


def _load_papers() -> list[dict]:
    ws = load_workbook(config.TRACKING_XLSX)["Papers"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    return [r for r in rows if not str(r["Paper ID"]).endswith("-copy")]


def _paper_text(pid: str, file_name: str) -> tuple[str, str]:
    path = config.PAPERS_DIR / file_name
    text = _full_text(path)
    if len(text) >= 200:
        return text, "text layer"
    cache = config.OCR_CACHE_DIR / f"{pid}.txt"
    log = lambda m: print(m, flush=True)
    return ocr_mod.ocr_pdf(path, OCR_MARKERS, cache_path=cache, log=log), "OCR (cached)"


def _write_sheets(question_rows: list[dict], checks: list[dict]) -> None:
    wb = load_workbook(config.TRACKING_XLSX)

    def existing(sheet: str) -> list[dict]:
        if sheet not in wb.sheetnames:
            return []
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]

    # merge: keep rows of papers not processed in this run
    seen_q = {r["Paper ID"] for r in question_rows}
    question_rows = [r for r in existing("Questions") if r["Paper ID"] not in seen_q] + question_rows
    seen_c = {r["Paper ID"] for r in checks}
    checks = [r for r in existing("Paper marks check") if r["Paper ID"] not in seen_c] + checks

    def qkey(r: dict):
        m = re.search(r"\d+", str(r["Q no."]))
        return (r["Paper ID"], int(m.group()) if m else 10**6, str(r["Q no."]))

    question_rows.sort(key=qkey)
    checks.sort(key=lambda r: r["Paper ID"])

    for sheet in ("Questions", "Paper marks check"):
        if sheet in wb.sheetnames:
            del wb[sheet]
    ws = wb.create_sheet("Questions")
    ws.append(QUESTIONS_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in question_rows:
        ws.append([_clean(row[h]) for h in QUESTIONS_HEADERS])
    widths = [16, 6, 8, 26, 20, 7, 12, 14, 18, 34, 80, 28]
    for col, width in zip("ABCDEFGHIJKL", widths):
        ws.column_dimensions[col].width = width
    dv = DataValidation(type="list", formula1='"' + ",".join(QUESTION_TYPES) + '"',
                        allow_blank=True)
    dv.add(f"E2:E{ws.max_row}")
    ws.add_data_validation(dv)

    wc = wb.create_sheet("Paper marks check")
    wc.append(CHECK_HEADERS)
    for cell in wc[1]:
        cell.font = Font(bold=True)
    for row in checks:
        wc.append([_clean(row[h]) for h in CHECK_HEADERS])
    for col, width in zip("ABCDE", (16, 14, 13, 42, 60)):
        wc.column_dimensions[col].width = width
    wb.save(config.TRACKING_XLSX)


def process_papers(papers: list[dict], tag: str) -> tuple[list[dict], list[dict]]:
    question_rows: list[dict] = []
    checks: list[dict] = []
    for paper in papers:
        pid = paper["Paper ID"]
        file_name = paper["File name"]
        official = paper["Official total marks"]
        official_int = official if isinstance(official, int) else None
        print(f"[{tag}] {pid} ({file_name})", flush=True)
        text, source = _paper_text(pid, file_name)
        conf = "LLM-proposed labels (unverified)" if source == "text layer" \
            else "LLM-proposed labels on OCR text (unverified)"
        blocks = _split_blocks(text)
        paper_norm = _norm(text)
        print(f"    {len(blocks)} blocks, labelling via LLM...", flush=True)
        # LLMs are unreliable with part labels - present blocks under globally
        # unique numbers and later route rows to blocks BY CONTENT, not labels.
        prompt_blocks: list[tuple[str, str]] = [
            (str(i), body) for i, (_, body) in enumerate(blocks, start=1)
        ]
        labelled = _label_paper(pid, prompt_blocks)

        from difflib import SequenceMatcher

        norm_blocks = [(q_no, _norm(body)) for q_no, body in blocks]
        assigned: dict[str, list[dict]] = {q_no: [] for q_no, _ in blocks}
        for row in labelled:
            qn = _norm(str(row["original_text"]))
            placed = False
            if qn:
                for q_no, nb in norm_blocks:  # exact/fuzzy quote -> that block
                    if qn in nb or _fuzzy_contains(qn, nb):
                        assigned[q_no].append(row)
                        placed = True
                        break
                if not placed:  # best-effort: longest match overall
                    best_q, best_ratio = None, 0.0
                    for q_no, nb in norm_blocks:
                        ratio = SequenceMatcher(
                            None, qn, nb).find_longest_match(
                            0, len(qn), 0, len(nb)).size / len(qn)
                        if ratio > best_ratio:
                            best_q, best_ratio = q_no, ratio
                    if best_ratio >= 0.5:
                        assigned[best_q].append(row)

        regex_total = 0.0
        known = 0.0
        any_unknown = False
        for q_no, block in blocks:
            bm = _block_marks(block)
            if bm is not None:
                regex_total += bm
            rows, ok = _verify_block(block, assigned.get(q_no, []),
                                     full_paper_norm=paper_norm, q_no=q_no)
            if not ok:
                print(f"    !! verbatim check failed for Q{q_no} - stored whole block",
                      flush=True)
            for row in rows:
                row["q_no"] = q_no  # the block's label is authoritative

                if row["marks"] is None:
                    any_unknown = True
                else:
                    known += float(row["marks"])
                question_rows.append({
                    "Paper ID": pid, "Q no.": row["q_no"], "Sub-part": row["sub_part"],
                    "Topic": row["topic"], "Question type": row["question_type"],
                    "Marks": row["marks"], "Has diagram": row["has_diagram"],
                    "Diagram type": row["diagram_type"],
                    "Machine-verifiable": row["machine_verifiable"],
                    "Wording note": row["wording_note"],
                    "Original question text": row["original_text"], "Confidence": conf,
                })
        if official_int is not None:
            if not any_unknown and abs(known - official_int) < 0.01:
                status, note = "match", f"logged {known} == official {official_int}"
            else:
                extra = " (+ some rows without printed marks)" if any_unknown else ""
                status = "mismatch"
                note = f"official {official_int} vs logged {known}{extra}"
        else:
            status = "official total unsure"
            extra = " (+ some rows without printed marks)" if any_unknown else ""
            note = f"computed sum = {known}{extra}"
        if abs(regex_total - known) >= 0.01:
            note += f"; regex-found marks total {regex_total}"
        checks.append({"Paper ID": pid, "Official total": official,
                       "Computed sum": known, "Status": status, "Note": note})
        print(f"    {len(blocks)} blocks, logged {known} marks ({status})", flush=True)
    return question_rows, checks


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", choices=["all", "text", "scanned"], default="all")
    ap.add_argument("--pids", default="",
                    help="comma-separated Paper IDs to process (ignores --only)")
    args = ap.parse_args()

    papers = _load_papers()
    if args.pids:
        wanted = {p.strip() for p in args.pids.split(",") if p.strip()}
        papers = [p for p in papers if p["Paper ID"] in wanted]
        missing = wanted - {p["Paper ID"] for p in papers}
        if missing:
            raise SystemExit(f"unknown Paper IDs: {sorted(missing)}")
    elif args.only == "text":
        papers = [p for p in papers if len(_full_text(config.PAPERS_DIR / p["File name"])) >= 200]
    elif args.only == "scanned":
        papers = [p for p in papers if len(_full_text(config.PAPERS_DIR / p["File name"])) < 200]
    print(f"logging questions for {len(papers)} papers ({args.only or args.pids})")
    q_rows, checks = process_papers(papers, args.only)
    _write_sheets(q_rows, checks)
    print(f"\nwrote {len(q_rows)} question rows and {len(checks)} check rows -> "
          f"{config.TRACKING_XLSX}")


if __name__ == "__main__":
    main()
