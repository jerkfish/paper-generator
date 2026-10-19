"""Recover whole-block fallback rows.

Whole-block rows (verbatim check failed in Step 3) are cleaned (footers,
page headers, page numbers stripped; whitespace rejoined), re-labelled by the
LLM and re-verified against the cleaned block + original paper text.
Recovered blocks are replaced by their sub-part rows in the Questions sheet;
unrecovered blocks stay flagged.

Run:  uv run python -m src.ingest.recover_blocks
"""
from __future__ import annotations

import re
from collections import defaultdict
from difflib import SequenceMatcher

from openpyxl import load_workbook

from src.config import TRACKING_XLSX
from src.fingerprint.types import is_whole_block_row
from src.ingest.log_questions import (_fuzzy_contains, _label_paper, _norm,
                                      _verify_block)

NOISE_LINE_PAT = re.compile(
    r"^\s*(CONFIDENTIAL|SULIT)\s*$"
    r"|^\s*CS[J/][A-Z0-9 ]{0,30}/\s*CSC662\s*$"
    r"|©\s*Hak\s*Cipta.*"
    r"|This examination paper consists of.*"
    r"|DO NOT TURN THIS PAGE.*"
    r"|===== PAGE \d+ ====="
    r"|^\s*\d{1,2}\s*$"
    r"|END OF QUESTIONS? PAPER.*", re.I)
INLINE_NOISE_PAT = re.compile(
    r"\bCONFIDENTIAL\b|CS[J/][A-Z0-9]{1,25}/CSC662|"
    r"[O@]\s*Hak Cipta Universiti Teknologi MARA", re.I)

SHEET_HEADERS = [
    "Paper ID", "Q no.", "Sub-part", "Topic", "Task", "Marks",
    "Has diagram", "Diagram type", "Machine-verifiable", "Wording note",
    "Original question text", "Confidence",
]


def clean_block(text: str) -> str:
    lines = [line for line in str(text).splitlines()
             if not NOISE_LINE_PAT.match(line)]
    joined = " ".join(lines)
    joined = INLINE_NOISE_PAT.sub(" ", joined)
    joined = re.sub(r"\s+", " ", joined)
    return joined.strip()


def _sheet_rows() -> list[dict]:
    ws = load_workbook(TRACKING_XLSX)["Questions"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def _route(blocks: list[tuple[str, str, str]], labelled: list[dict]) -> dict[str, list[dict]]:
    """Content-based routing of LLM rows to cleaned blocks."""
    assigned: dict[str, list[dict]] = {q: [] for q, _, _ in blocks}
    for row in labelled:
        qn = _norm(str(row["original_text"]))
        if not qn:
            continue
        for q_no, _, clean in blocks:
            nc = _norm(clean)
            if qn in nc or _fuzzy_contains(qn, nc):
                assigned[q_no].append(row)
                break
        else:
            best_q, best_ratio = None, 0.0
            for q_no, _, clean in blocks:
                nc = _norm(clean)
                ratio = SequenceMatcher(None, qn, nc).find_longest_match(
                    0, len(qn), 0, len(nc)).size / len(qn)
                if ratio > best_ratio:
                    best_q, best_ratio = q_no, ratio
            if best_ratio >= 0.5:
                assigned[best_q].append(row)
    return assigned


def _new_row(template: dict, q_no: str, lab: dict, flag: str) -> dict:
    return {
        "Paper ID": template["Paper ID"], "Q no.": q_no,
        "Sub-part": lab.get("sub_part", ""),
        "Topic": "", "Task": "",
        "Marks": lab.get("marks"),
        "Has diagram": lab.get("has_diagram", False),
        "Diagram type": lab.get("diagram_type", ""),
        "Machine-verifiable": "",
        "Wording note": (lab.get("wording_note") or "") + f"; {flag}",
        "Original question text": lab.get("original_text", ""),
        "Confidence": template.get("Confidence", ""),
    }


def main() -> None:
    rows = _sheet_rows()
    fallback = [r for r in rows if is_whole_block_row(r)]

    # original paper text (for "quote matched elsewhere" verification)
    from src.config import OCR_CACHE_DIR, PAPERS_DIR
    from src.ingest.log_questions import _full_text

    pwb = load_workbook(TRACKING_XLSX)["Papers"]
    papers = {r["Paper ID"]: r["File name"] for r in
              (dict(zip([c.value for c in pwb[1]], row))
               for row in pwb.iter_rows(min_row=2, values_only=True))}
    paper_texts: dict[str, str] = {}
    for pid in {r["Paper ID"] for r in fallback}:
        text = _full_text(PAPERS_DIR / papers[pid])
        if len(text) < 200:
            cache = OCR_CACHE_DIR / f"{pid}.txt"
            if cache.exists():
                text = cache.read_text(encoding="utf-8")
        paper_texts[pid] = text

    recovered: dict[tuple[str, str], list[dict]] = {}
    failed: list[tuple[str, str, str]] = []

    by_paper: dict[str, list[dict]] = defaultdict(list)
    for r in fallback:
        by_paper[r["Paper ID"]].append(r)

    for pid, frs in by_paper.items():
        blocks = [(r["Q no."], r["Original question text"],
                   clean_block(r["Original question text"])) for r in frs]
        print(f"[recover] {pid}: {len(blocks)} blocks", flush=True)
        labelled = _label_paper(pid, [(q, c) for q, _, c in blocks])
        assigned = _route(blocks, labelled)
        paper_norm = _norm(paper_texts.get(pid, ""))
        for q_no, raw, clean in blocks:
            rows_l, ok = _verify_block(clean, assigned.get(q_no, []),
                                       full_paper_norm=paper_norm, q_no=q_no)
            if ok and rows_l:
                template = next(r for r in frs if r["Q no."] == q_no)
                recovered[(pid, q_no)] = [
                    _new_row(template, q_no, row, "recovered by cleaning pass")
                    for row in rows_l]
            else:
                failed.append((pid, q_no, clean[:90]))

    # write back
    out: list[dict] = []
    for r in rows:
        key = (r["Paper ID"], r["Q no."])
        if is_whole_block_row(r) and key in recovered:
            out.extend(recovered[key])
        else:
            out.append(r)

    wb = load_workbook(TRACKING_XLSX)
    ws = wb["Questions"]
    ws.delete_rows(2, ws.max_row)
    for row in out:
        ws.append([row.get(h) for h in SHEET_HEADERS])
    wb.save(TRACKING_XLSX)

    rec_marks = sum(float(r["Marks"]) for rr in recovered.values()
                    for r in rr if r["Marks"] is not None)
    fb_marks = sum(float(r["Marks"]) for r in fallback if r["Marks"] is not None)
    print(f"\nrecovered {len(recovered)} blocks ({sum(len(rr) for rr in recovered.values())} "
          f"rows, {rec_marks:.0f} marks of {fb_marks:.0f})")
    print(f"still failed: {len(failed)}")
    for pid, q, snip in failed:
        print(f"  {pid} Q{q}: {snip}")


if __name__ == "__main__":
    main()
