"""Step 2 - Register papers.

Scans data/papers/, extracts registration fields from each PDF and fills the
"Papers" sheet in data/tracking/mock_paper_planner.xlsx.

Field extraction order: PDF text layer first; if the cover has no text layer
(scanned paper), the cover scan is OCR'd with the built-in Windows OCR engine
(WinRT - no extra installs). Values that cannot be extracted with confidence
are recorded as "unsure" and listed in the Notes column - nothing is guessed.

Run:  uv run python -m src.ingest.register_papers
"""
from __future__ import annotations

import hashlib
import pathlib
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.config import OCR_MIN_TEXT_CHARS, PAPERS_DIR, TRACKING_XLSX

MONTH_ORDER = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
    "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]

FILE_PAT = re.compile(r"CSC662\s+([A-Z]+)\s+(\d{4})(?:\s+\(2\))?\.pdf$", re.I)
# TEKNOL\w* tolerates text-layer glyph mangling ("TEKNOLQGl" for "TEKNOLOGI")
UNIVERSITY_PAT = re.compile(r"UNIVERSITI\s+TEKNOL\w*\s+MARA", re.I)
COURSE_PAT = re.compile(r"\bCSC\s*662\b", re.I)
DURATION_PAT = re.compile(r"(\d+)\s*HOURS?", re.I)
YEAR_PAT = re.compile(r"(?<!\d)(?:19|20)(\d{2})(?!\d)")

HEADERS = [
    "Paper ID", "File name", "Month", "Year", "University",
    "Course code", "Era", "Official total marks", "Duration", "Notes",
]

# Windows built-in OCR (WinRT) invoked via PowerShell. Takes one image path.
_WIN_OCR_PS1 = r"""
param([string]$ImagePath)
$ErrorActionPreference = "Stop"
Add-Type -AssemblyName System.Runtime.WindowsRuntime | Out-Null
[Windows.Storage.StorageFile,Windows.Storage,ContentType=WindowsRuntime] | Out-Null
[Windows.Storage.Streams.IRandomAccessStream,Windows.Storage.Streams,ContentType=WindowsRuntime] | Out-Null
[Windows.Media.Ocr.OcrEngine,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Media.Ocr.OcrResult,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Graphics.Imaging.BitmapDecoder,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Graphics.Imaging.SoftwareBitmap,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
[Windows.Globalization.Language,Windows.Foundation,ContentType=WindowsRuntime] | Out-Null
$asTaskGeneric = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Await($WinRtTask, $ResultType) {
    $asTask = $asTaskGeneric.MakeGenericMethod($ResultType)
    $netTask = $asTask.Invoke($null, @($WinRtTask))
    $netTask.Wait(-1) | Out-Null
    $netTask.Result
}
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
if ($null -eq $engine) {
    $lang = New-Object Windows.Globalization.Language "en-US"
    $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage($lang)
}
if ($null -eq $engine) { Write-Output "OCR_ENGINE_UNAVAILABLE"; exit 1 }
$file = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync($ImagePath)) ([Windows.Storage.StorageFile])
$stream = Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
$decoder = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
$bitmap = Await ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
$result = Await ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
Write-Output $result.Text
"""


def _md5(path: pathlib.Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def _full_text(path: pathlib.Path) -> str:
    """All extractable text in the PDF, page by page."""
    from pypdf import PdfReader

    parts = []
    try:
        reader = PdfReader(str(path))
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                pass
    except Exception:
        return ""
    return "\n".join(parts)


def _win_ocr(image_path: str) -> str:
    """OCR one image with the Windows engine. Returns "" if unavailable."""
    if not sys.platform.startswith("win"):
        return ""
    with tempfile.TemporaryDirectory() as td:
        ps1 = pathlib.Path(td) / "ocr.ps1"
        ps1.write_text(_WIN_OCR_PS1, encoding="utf-8")
        try:
            proc = subprocess.run(
                ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
                 "-File", str(ps1), image_path],
                capture_output=True, timeout=120,
            )
        except Exception:
            return ""
        # PowerShell console output is not UTF-8 - decode leniently.
        return proc.stdout.decode("utf-8", errors="replace").strip().lstrip("﻿")


def _ocr_cover_text(path: pathlib.Path) -> str:
    """OCR the cover scan (trying rotations; stitching strip-scans). "" on failure."""
    from PIL import Image
    from pypdf import PdfReader

    try:
        page = PdfReader(str(path)).pages[0]
        imgs = [im.image for im in page.images if im.image is not None]
    except Exception:
        return ""
    if not imgs:
        return ""

    variants: list[Image.Image] = list(imgs)
    if len(imgs) > 1 and len({im.width for im in imgs}) <= 2:
        # some scanners store one page as horizontal strips - stitch them
        width = max(im.width for im in imgs)
        height = sum(im.height for im in imgs)
        stitched = Image.new("RGB", (width, height), "white")
        y = 0
        for im in imgs:
            stitched.paste(im.convert("RGB"), (0, y))
            y += im.height
        variants.append(stitched)

    def score(text: str) -> int:
        """Prefer long text that contains expected cover markers (right way up)."""
        if not text:
            return 0
        s = len(text)
        if UNIVERSITY_PAT.search(text):
            s += 300
        if COURSE_PAT.search(text):
            s += 300
        if DURATION_PAT.search(text):
            s += 200
        if re.search(r"PART\s*A\b", text, re.I):
            s += 100
        return s

    best_text, best_score = "", 0
    with tempfile.TemporaryDirectory() as td:
        for i, im in enumerate(variants):
            base = pathlib.Path(td) / f"v{i}.png"
            im.convert("RGB").save(base)
            angles = [0]
            text = _win_ocr(str(base))
            s = score(text)
            if s > best_score:
                best_text, best_score = text, s
            if s < 400:
                angles += [90, 180, 270]  # cover stored rotated (seen in the wild)
            for angle in angles[1:]:
                rp = pathlib.Path(td) / f"v{i}-r{angle}.png"
                Image.open(base).rotate(angle, expand=True).save(rp)
                text = _win_ocr(str(rp))
                s = score(text)
                if s > best_score:
                    best_text, best_score = text, s
    return best_text


def _printed_total(cover: str) -> int | None:
    """An official total printed on the cover, e.g. 'TOTAL MARKS /80 /20 /100'."""
    candidates = []
    lines = cover.splitlines()
    for idx, line in enumerate(lines):
        mt = re.search(r"\bTOTAL\b", line, re.I)
        if not mt:
            continue
        window = line[mt.start():] + " "
        if re.search(r"MARKS?", window, re.I) and idx + 1 < len(lines):
            nxt = lines[idx + 1]
            # grid values ("/80 /20 /100") may spill onto the next line
            if re.fullmatch(r"[\s/0-9]*", nxt) and re.search(r"\d", nxt):
                window += " " + nxt
        nums = [int(n) for n in re.findall(r"(?<!\d)(\d{2,3})(?!\d)", window)]
        if nums and 10 <= nums[-1] <= 200:
            candidates.append(nums[-1])
    uniq = set(candidates)
    return uniq.pop() if len(uniq) == 1 else None


def _extract_row(path: pathlib.Path, dup_map: dict[str, str]) -> dict:
    m = FILE_PAT.match(path.name)
    month = m.group(1).upper() if m else "unsure"
    year = int(m.group(2)) if m else "unsure"
    is_copy = "(2)" in path.name

    full_text = _full_text(path)
    cover = full_text[:4000]

    notes = []
    if len(cover.strip()) < OCR_MIN_TEXT_CHARS:
        ocr = _ocr_cover_text(path)
        if len(ocr) >= OCR_MIN_TEXT_CHARS:
            cover = cover + "\n" + ocr
            notes.append("no text layer (scanned); cover read via Windows OCR")
        else:
            notes.append("no text layer (scanned); cover OCR failed")

    university = "Universiti Teknologi MARA" if UNIVERSITY_PAT.search(cover) else "unsure"
    course = "CSC662" if COURSE_PAT.search(cover) else "unsure"
    dur = DURATION_PAT.search(cover)
    duration = f"{dur.group(1)} hours" if dur else "unsure"

    text_years = {2000 + int(y) for y in YEAR_PAT.findall(cover)}
    if isinstance(year, int) and text_years and year not in text_years:
        notes.append(f"year {year} taken from file name; not confirmed in cover text")

    # Era is a provisional YEAR-based bucket (visible format shifts happen
    # around 2013 - new cover template - and 2018 - standard Part A/B).
    if isinstance(year, int):
        if year <= 2012:
            era = "classic (provisional)"
        elif year <= 2017:
            era = "transitional (provisional)"
        else:
            era = "modern (provisional)"
        notes.append("era provisional (year-based bucket)")
    else:
        era = "unsure"
        notes.append("unsure: era")

    has_parts = bool(re.search(r"PART\s*A\b", cover, re.I))
    single_list = bool(re.search(r"consists?\s+of", cover, re.I))
    if has_parts:
        notes.append("cover has Part A/B structure")
        if isinstance(year, int) and year <= 2012:
            notes.append("Part A/B structure is unusual for this era")
    elif single_list:
        notes.append("cover has single-list structure")
        if isinstance(year, int) and year >= 2018:
            notes.append("single-list structure (most 2018+ papers use Part A/B)")

    total = _printed_total(cover)
    total_field = total if total is not None else "unsure"
    if total is not None:
        notes.append(f"printed total {total} marks found on cover")

    if is_copy:
        canonical = dup_map.get(path.name)
        if canonical:
            notes.append(f"byte-identical copy of {canonical}")
        else:
            notes.append("marked '(2)' but no byte-identical twin found")

    unsure = [
        name for name, value in [
            ("university", university), ("course code", course),
            ("duration", duration), ("official total marks", total_field),
            ("era", era),
        ] if value == "unsure"
    ]
    if unsure:
        notes.append("unsure: " + ", ".join(unsure))

    pid = (f"CSC662-{year}-{MONTH_ORDER.index(month) + 1:02d}"
           if isinstance(year, int) and month in MONTH_ORDER else "unsure")
    if is_copy:
        pid += "-copy"

    return {
        "Paper ID": pid, "File name": path.name, "Month": month, "Year": year,
        "University": university, "Course code": course, "Era": era,
        "Official total marks": total_field, "Duration": duration,
        "Notes": "; ".join(notes) if notes else "-",
    }


def _write_sheet(rows: list[dict]) -> None:
    TRACKING_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(TRACKING_XLSX) if TRACKING_XLSX.exists() else Workbook()
    if "Papers" in wb.sheetnames:
        del wb["Papers"]
    ws = wb.create_sheet("Papers", 0)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[h] for h in HEADERS])
    for col, width in zip("ABCDEFGHIJ", (18, 30, 11, 7, 28, 11, 24, 20, 10, 70)):
        ws.column_dimensions[col].width = width
    wb.save(TRACKING_XLSX)


def _print_summary(rows: list[dict]) -> None:
    print(f"Registered {len(rows)} papers -> {TRACKING_XLSX}")
    for r in rows:
        print(f"  {r['Paper ID']:22} {r['Month']:9} {r['Year']}  era={r['Era']:24} "
              f"total={r['Official total marks']}  dur={r['Duration']}")
    print("\nUnsure / flags:")
    for r in rows:
        if r["Notes"] != "-":
            print(f"  {r['File name']}: {r['Notes']}")


def main() -> None:
    files = sorted(p for p in PAPERS_DIR.iterdir() if p.suffix.lower() == ".pdf")
    if not files:
        raise SystemExit(f"no PDFs found in {PAPERS_DIR}")

    by_hash: dict[str, list[str]] = {}
    for p in files:
        by_hash.setdefault(_md5(p), []).append(p.name)
    dup_map = {}
    for names in by_hash.values():
        if len(names) > 1:
            # prefer the non-"(2)" file as canonical
            canonical = sorted(names, key=lambda n: ("(2)" in n, n))[0]
            for n in names:
                if n != canonical:
                    dup_map[n] = canonical

    rows = [_extract_row(p, dup_map) for p in files]

    def sort_key(r: dict):
        year = r["Year"] if isinstance(r["Year"], int) else 9999
        month = MONTH_ORDER.index(r["Month"]) if r["Month"] in MONTH_ORDER else 99
        return (year, month, r["Paper ID"])

    rows.sort(key=sort_key)
    _write_sheet(rows)
    _print_summary(rows)


if __name__ == "__main__":
    main()
