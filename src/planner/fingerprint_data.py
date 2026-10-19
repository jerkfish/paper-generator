"""Modern-fingerprint inputs for the planner.

Reads the Questions sheet, keeps the 10 clean modern papers (year >= 2018 and
marks sum 98..102), and exposes topic/task/marks weights plus the per-paper
machine-verifiable (computed) mark distribution.

`risk/management` is normalised to `risk/management (ethics)` so the ethics
deck can be attached to it.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from functools import lru_cache

from openpyxl import load_workbook

from src.config import TRACKING_XLSX
from src.fingerprint.types import is_whole_block_row

CLEAN_RANGE = (98, 102)
THEORY_TOPICS = [
    "attacks/malware",
    "authentication/biometrics",
    "access control models",
    "trusted computing/blockchain",
    "software/OS security",
    "network security",
    "risk/management (ethics)",
]
TOPIC_ALIASES = {"risk/management": "risk/management (ethics)"}
# Modern non-compute marks that fall outside the seven theory topics.  They
# are reported as coverage gaps rather than silently reallocated.
EXCLUDED_TOPICS = {
    "other", "general cryptography", "classical ciphers", "public-key",
    "hashing/signatures", "block ciphers",
}


def normalize_topic(topic: str) -> str:
    return TOPIC_ALIASES.get(str(topic), str(topic))


@dataclass
class ModernStats:
    clean_ids: list[str]
    paper_marks: dict[str, int]
    mv_per_paper: dict[str, int]
    mv_distribution: list[int]
    topic_marks: dict[str, int]
    task_marks: dict[tuple[str, str], int] = field(default_factory=dict)
    rows: list[dict] = field(default_factory=list)
    excluded_topic_marks: dict[str, int] = field(default_factory=dict)

    @property
    def theory_marks(self) -> int:
        return sum(self.topic_marks.values())

    @property
    def topic_weights(self) -> dict[str, float]:
        total = sum(self.topic_marks.values()) or 1
        return {t: self.topic_marks.get(t, 0) / total for t in THEORY_TOPICS}

    def rows_for(self, topic: str, marks: int) -> list[dict]:
        topic = normalize_topic(topic)
        return [r for r in self.rows
                if normalize_topic(r["Topic"]) == topic
                and int(r["Marks"] or 0) == marks
                and str(r["Task"]) != "compute"]

    def task_weights(self, topic: str) -> dict[str, int]:
        topic = normalize_topic(topic)
        out = defaultdict(int)
        for (t, task), marks in self.task_marks.items():
            if normalize_topic(t) == topic:
                out[task] += marks
        return dict(out)


def _question_rows() -> list[dict]:
    wb = load_workbook(TRACKING_XLSX, read_only=False, data_only=True)
    ws = wb["Questions"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return [r for r in rows if not is_whole_block_row(r)]


@lru_cache(maxsize=1)
def load_all_rows() -> tuple[dict, ...]:
    """All 35 papers, with the topic alias applied; compute rows retained."""
    return tuple(dict(r, Topic=normalize_topic(r["Topic"]))
                 for r in _question_rows())


@lru_cache(maxsize=1)
def load_modern_stats() -> ModernStats:
    rows = load_all_rows()

    paper_marks: dict[str, int] = defaultdict(int)
    mv: dict[str, int] = defaultdict(int)
    year: dict[str, int] = {}
    for r in rows:
        pid = str(r["Paper ID"])
        marks = int(r["Marks"] or 0)
        paper_marks[pid] += marks
        year[pid] = int(pid.split("-")[1])
        if str(r["Machine-verifiable"]).upper() == "Y":
            mv[pid] += marks

    clean_ids = sorted(pid for pid in paper_marks
                       if year[pid] >= 2018
                       and CLEAN_RANGE[0] <= paper_marks[pid] <= CLEAN_RANGE[1])
    clean_set = set(clean_ids)

    topic_marks: dict[str, int] = defaultdict(int)
    task_marks: dict[tuple[str, str], int] = defaultdict(int)
    excluded: dict[str, int] = defaultdict(int)
    clean_rows: list[dict] = []
    for r in rows:
        if str(r["Paper ID"]) not in clean_set:
            continue
        topic = normalize_topic(r["Topic"])
        task = str(r["Task"])
        marks = int(r["Marks"] or 0)
        clean_rows.append(r)
        if task == "compute":
            continue  # computed slots are handled separately
        if topic not in THEORY_TOPICS:
            excluded[topic] += marks
            continue
        topic_marks[topic] += marks
        task_marks[(topic, task)] += marks

    return ModernStats(
        clean_ids=clean_ids,
        paper_marks={p: paper_marks[p] for p in clean_ids},
        mv_per_paper={p: mv[p] for p in clean_ids},
        mv_distribution=sorted(mv[p] for p in clean_ids),
        topic_marks=dict(topic_marks),
        task_marks=dict(task_marks),
        rows=clean_rows,
        excluded_topic_marks=dict(excluded),
    )


@lru_cache(maxsize=1)
def load_part_b_stats() -> dict[str, dict]:
    """Part B topic and command-word marks from the 8 clean A/B papers.

    Kept separate from the overall modern weights because Part B is a
    different task mix (the brief asks Part B tasks not to inherit Part A's).
    """
    import re

    from src.planner.tasks import command_word

    rows = load_all_rows()
    stats = load_modern_stats()
    clean = set(stats.clean_ids)
    topics: dict[str, int] = defaultdict(int)
    labels: dict[str, int] = defaultdict(int)
    labels_at_marks: dict[tuple[int, str], int] = defaultdict(int)
    for r in rows:
        if str(r["Paper ID"]) not in clean:
            continue
        if not re.match(r"^B\d+$", str(r["Q no."]).strip()):
            continue
        marks = int(r["Marks"] or 0)
        topics[normalize_topic(r["Topic"])] += marks
        label, _ = command_word(r["Original question text"])
        if label != "Other":
            labels[label] += marks
            labels_at_marks[(marks, label)] += marks
    return {
        "topic_marks": dict(topics),
        "label_marks": dict(labels),
        "label_marks_at_marks": {f"{m}:{l}": v
                                 for (m, l), v in labels_at_marks.items()},
    }


if __name__ == "__main__":
    stats = load_modern_stats()
    print("clean modern papers:", len(stats.clean_ids))
    print("marks total:", sum(stats.paper_marks.values()))
    print("MV distribution:", stats.mv_distribution,
          "mean", sum(stats.mv_distribution) / len(stats.mv_distribution))
    print("topic marks:", stats.topic_marks)
    print("topic weights:", {k: round(v, 3) for k, v in stats.topic_weights.items()})
    print("excluded topic marks:", stats.excluded_topic_marks)
