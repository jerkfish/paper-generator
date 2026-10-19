"""Read-only lecture-deck ingestion, chunked by deck + slide.

The decks live outside the project (KB folder path in config).  Nothing is
written there; the extracted chunks are cached under data/tracking/theory/.
"""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from pptx import Presentation

from src import config

MAX_CHUNK_CHARS = 1200
MIN_CHUNK_CHARS = 40


@dataclass
class DeckChunk:
    chunk_id: str
    deck: str
    slide: int
    part: int
    text: str
    chars: int


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False):
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip() or p.text.strip()
            if t:
                out.append(t)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
        for sub in shape.shapes:
            out.extend(_shape_text(sub))
    return out


def _slide_text(slide) -> str:
    texts: list[str] = []
    for shape in slide.shapes:
        texts.extend(_shape_text(shape))
    # De-duplicate consecutive identical lines (decorative text boxes repeat).
    clean: list[str] = []
    for t in texts:
        if not clean or clean[-1] != t:
            clean.append(t)
    return "\n".join(clean)


def _split_text(text: str, limit: int = MAX_CHUNK_CHARS) -> list[str]:
    if len(text) <= limit:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        if current and len(current) + 1 + len(sentence) > limit:
            chunks.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        chunks.append(current)
    return chunks


def _source_fingerprint(folder: Path) -> list[dict]:
    out = []
    for f in sorted(folder.glob("*.pptx")):
        st = f.stat()
        out.append({"name": f.name, "size": st.st_size,
                    "mtime_ns": st.st_mtime_ns})
    return out


def _cache_path() -> Path:
    return config.THEORY_CACHE_DIR / "deck_chunks.json"


def extract_chunks(folder: Path) -> list[DeckChunk]:
    if not folder.exists():
        raise RuntimeError(f"lecture deck folder does not exist: {folder}")
    files = sorted(folder.glob("*.pptx"))
    if not files:
        raise RuntimeError(f"no .pptx decks found in {folder}")
    chunks: list[DeckChunk] = []
    for f in files:
        try:
            prs = Presentation(str(f))
        except Exception as e:  # keep going; report the unreadable deck
            chunks.append(DeckChunk(
                chunk_id=f"{f.stem}::ERROR::0", deck=f.stem, slide=0,
                part=0, text=f"UNREADABLE DECK: {e}", chars=0))
            continue
        for slide_no, slide in enumerate(prs.slides, 1):
            text = _slide_text(slide)
            if len(text) < MIN_CHUNK_CHARS:
                continue
            for part, piece in enumerate(_split_text(text), 1):
                if len(piece) < MIN_CHUNK_CHARS:
                    continue
                chunks.append(DeckChunk(
                    chunk_id=f"{f.stem}::s{slide_no:02d}::{part}",
                    deck=f.stem, slide=slide_no, part=part, text=piece,
                    chars=len(piece)))
    return chunks


def load_chunks(refresh: bool = False) -> list[DeckChunk]:
    folder = Path(config.LECTURE_DECKS_DIR)
    fingerprint = _source_fingerprint(folder) if folder.exists() else []
    cache = _cache_path()
    if not refresh and cache.exists():
        try:
            payload = json.loads(cache.read_text(encoding="utf-8"))
            if payload.get("source_fingerprint") == fingerprint:
                return [DeckChunk(**c) for c in payload["chunks"]]
        except Exception:
            pass
    chunks = extract_chunks(folder)
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps({
        "source_dir": str(folder),
        "source_fingerprint": fingerprint,
        "chunks": [asdict(c) for c in chunks],
    }, indent=2), encoding="utf-8")
    return chunks


def coverage(chunks: list[DeckChunk] | None = None) -> dict:
    chunks = chunks if chunks is not None else load_chunks()
    by_deck: dict[str, dict] = {}
    for c in chunks:
        d = by_deck.setdefault(c.deck, {"slides": set(), "chunks": 0,
                                        "chars": 0})
        d["slides"].add(c.slide)
        d["chunks"] += 1
        d["chars"] += c.chars
    out = {}
    for deck, d in sorted(by_deck.items()):
        out[deck] = {"slides_with_text": len(d["slides"]),
                     "chunks": d["chunks"], "chars": d["chars"]}
    return out


if __name__ == "__main__":
    cs = load_chunks()
    print(f"chunks: {len(cs)}")
    cov = coverage(cs)
    for deck, d in cov.items():
        print(f"{deck}: {d['slides_with_text']} slides, {d['chunks']} chunks")
    ethics = [c for c in cs if c.deck.startswith("071")]
    print(f"\n071 Ethics: {len(ethics)} chunks, slides "
          f"{sorted({c.slide for c in ethics})}")
