"""Web job lifecycle, cost cap, cache and Library tests."""
from __future__ import annotations

import json
import time
from pathlib import Path

from src.app import jobs
from src.theory.generate import CostCapExceeded, GenerationCancelled


def _wait(job, timeout=5.0):
    end = time.time() + timeout
    while job.running and time.time() < end:
        time.sleep(0.02)
    assert not job.running, f"job stuck in {job.status}: {job.error}"


def _fake_env(monkeypatch, tmp_path):
    monkeypatch.setattr(jobs.config, "WEB_RUNS_DIR", tmp_path / "runs")
    monkeypatch.setattr(jobs.config, "WEB_CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(jobs.config, "DEEPSEEK_API_KEY", "fake-key")
    monkeypatch.setattr(jobs.config, "PUBLIC_BUILD", False)
    monkeypatch.setattr(
        jobs, "build_plan",
        lambda mode, style, cm, seed: ([], {
            "seed": seed, "computed_marks": 0,
            "assigned_theory_topic_marks": {}, "slots": []}))
    monkeypatch.setattr(jobs, "to_markdown", lambda summary: "# plan")
    monkeypatch.setattr(jobs, "theory_report", lambda payload: "# theory")


def _fake_payload():
    return {
        "model": "fake-model",
        "accepted": [{"slot": "A1a", "marks": 2, "topic": "x",
                      "question": "Q?"}],
        "placeholders": [],
        "usage": {"cost_usd": 0.0123, "model": "fake-model", "calls": 2},
        "groups": [],
    }


def test_job_lifecycle(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)

    def fake_generate(plan_path, **kwargs):
        kwargs["progress_cb"]({"question": "A1", "slot": "A1a",
                               "status": "accepted", "cost_usd": 0.0123,
                               "accepted": 1, "placeholders": 0})
        return _fake_payload()

    def fake_paper(plan, theory, docx, pdf, public=False):
        Path(docx).write_bytes(b"docx")
        Path(pdf).write_bytes(b"pdf")
        return {"docx": str(docx), "pdf": str(pdf), "pdf_error": ""}

    def fake_scheme(paper, plan, theory, docx, pdf, public=False):
        Path(docx).write_bytes(b"docx")
        Path(pdf).write_bytes(b"pdf")
        return {"docx": str(docx), "pdf": str(pdf), "pdf_error": ""}

    monkeypatch.setattr(jobs, "generate", fake_generate)
    monkeypatch.setattr(jobs, "build_paper", fake_paper)
    monkeypatch.setattr(jobs, "build_scheme", fake_scheme)
    manager = jobs.JobManager()
    job = manager.create("realistic", 99, None)
    _wait(job)
    assert job.status == "done", job.error
    assert job.accepted == 1
    assert job.cost_usd == 0.0123
    assert Path(job.files["paper_pdf"]).exists()
    assert job.summary["placeholders"] == 0


def test_cost_cap_abort(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)

    def fake_generate(*args, **kwargs):
        raise CostCapExceeded(0.06, 0.05)

    monkeypatch.setattr(jobs, "generate", fake_generate)
    manager = jobs.JobManager()
    job = manager.create("realistic", 98, None)
    _wait(job)
    assert job.status == "cost_capped"
    assert "cost cap" in job.message.lower()
    assert job.cost_usd == 0.06


def test_cancel_running_job(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)

    def fake_generate(*args, **kwargs):
        for _ in range(200):
            if kwargs.get("should_stop") and kwargs["should_stop"]():
                raise GenerationCancelled("job cancelled")
            time.sleep(0.01)
        return _fake_payload()

    monkeypatch.setattr(jobs, "generate", fake_generate)
    manager = jobs.JobManager()
    job = manager.create("realistic", 97, None)
    time.sleep(0.05)
    assert manager.cancel(job.id)
    _wait(job)
    assert job.status == "cancelled"


def test_cache_hit_needs_no_key(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)
    monkeypatch.setattr(jobs.config, "DEEPSEEK_API_KEY", "")
    out = jobs.cache_path("realistic", 7, None, public=False)
    out.mkdir(parents=True)
    (out / "paper.docx").write_bytes(b"docx")
    (out / "meta.json").write_text(json.dumps({
        "seed": 7, "mode": "realistic", "computed_marks": None,
        "model": "fake-model", "cost_usd": 0.02, "accepted": 1,
        "placeholders": 0, "files": {"paper_docx": str(out / "paper.docx")},
        "summary": {"accepted": 1, "placeholders": 0},
        "progress": [{"question": "A1", "slot": "A1a",
                      "status": "accepted"}],
    }), encoding="utf-8")
    called = {"n": 0}

    def fake_generate(*args, **kwargs):
        called["n"] += 1
        return _fake_payload()

    monkeypatch.setattr(jobs, "generate", fake_generate)
    manager = jobs.JobManager()
    job = manager.create("realistic", 7, None)
    assert job.status == "done"
    assert job.cached is True
    assert job.cost_usd == 0.0
    assert called["n"] == 0


def test_generate_cost_cap_stops_after_call(monkeypatch):
    from src.llm import LLMResult
    from src.theory import generate as gen_mod
    plan = Path("outputs/eval_final_deepseek/seed_12_planner.json")

    def fake_chat(system, user, model=None, temperature=0.2, usage=None):
        usage.add(1_000_000, 0, 1_000_000, model="deepseek-flash")
        return LLMResult(data={"questions": []}, model="deepseek-flash",
                         usage=usage)

    monkeypatch.setattr(gen_mod, "chat_json_result", fake_chat)
    try:
        gen_mod.generate(plan, limit=1, cost_cap_usd=0.01)
    except CostCapExceeded as e:
        assert e.cost_usd >= 0.01
    else:
        raise AssertionError("cost cap did not fire")


def test_public_build_disables_library(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)
    monkeypatch.setattr(jobs.config, "PUBLIC_BUILD", True)
    assert jobs.library_items() == []
    assert jobs.library_file(1, "paper_pdf") is None


def test_publicise_replaces_branding(tmp_path):
    import docx
    from src.assemble.generic_build import publicise
    path = tmp_path / "branded.docx"
    d = docx.Document()
    d.add_paragraph("UNIVERSITI TEKNOLOGI MARA")
    d.add_paragraph("CS/MOCK OCT 2026/CSC662")
    d.save(str(path))
    publicise(path)
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    assert "GENERIC UNIVERSITY" in text
    assert "MOCK EXAM 2026" in text
    assert "UNIVERSITI TEKNOLOGI MARA" not in text


def test_library_works_without_key(monkeypatch, tmp_path):
    _fake_env(monkeypatch, tmp_path)
    monkeypatch.setattr(jobs.config, "DEEPSEEK_API_KEY", "")
    src = tmp_path / "eval_final_deepseek"
    src.mkdir(parents=True)
    (src / "seed_07_theory.json").write_text("{}", encoding="utf-8")
    monkeypatch.setattr(jobs.config, "OUTPUTS_DIR", tmp_path)
    out = jobs.cache_path("realistic", 7, None, public=False)
    out.mkdir(parents=True)
    paper = out / "paper.pdf"
    paper.write_bytes(b"pdf")
    (out / "meta.json").write_text(json.dumps({
        "seed": 7, "mode": "realistic", "model": "fake-model",
        "cost_usd": 0.0, "accepted": 1, "placeholders": 0,
        "files": {"paper_pdf": str(paper)}, "summary": {},
        "progress": [],
    }), encoding="utf-8")
    items = jobs.library_items()
    assert len(items) == 1
    assert items[0]["seed"] == 7
    assert jobs.library_file(7, "paper_pdf") == paper
