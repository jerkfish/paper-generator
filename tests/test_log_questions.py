"""Step 3 done-when tests: marks per paper match the official total or the
mismatch is listed; question rows keep the original wording."""
from __future__ import annotations

import re

from openpyxl import load_workbook

from src.config import TRACKING_XLSX

QUESTIONS_HEADERS = [
    "Paper ID", "Q no.", "Sub-part", "Topic", "Question type", "Marks",
    "Has diagram", "Diagram type", "Machine-verifiable", "Wording note",
    "Original question text", "Confidence",
]


def _sheet(name: str) -> list[dict]:
    ws = load_workbook(TRACKING_XLSX)[name]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def _papers() -> list[dict]:
    ws = load_workbook(TRACKING_XLSX)["Papers"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    return [r for r in rows if not str(r["Paper ID"]).endswith("-copy")]


def test_questions_sheet_has_expected_columns():
    ws = load_workbook(TRACKING_XLSX)["Questions"]
    assert [c.value for c in ws[1]] == QUESTIONS_HEADERS


def test_every_paper_has_question_rows_and_a_check_row():
    qrows = _sheet("Questions")
    checks = _sheet("Paper marks check")
    pids = {r["Paper ID"] for r in _papers()}
    assert {r["Paper ID"] for r in qrows} == pids
    assert {r["Paper ID"] for r in checks} == pids


def test_marks_values_sane():
    for r in _sheet("Questions"):
        if r["Marks"] is not None:
            assert isinstance(r["Marks"], (int, float)) and 0 < r["Marks"] <= 40, r


def test_controlled_vocabulary():
    for r in _sheet("Questions"):
        assert r["Machine-verifiable"] in {"Y", "Partial", "N"}
        assert r["Original question text"].strip()


def test_feb2022_marks_match_official_total():
    checks = {r["Paper ID"]: r for r in _sheet("Paper marks check")}
    row = checks["CSC662-2022-02"]
    assert row["Official total"] == 100
    assert row["Status"] == "match", row


def test_mismatches_are_listed():
    for r in _sheet("Paper marks check"):
        if r["Status"] not in {"match", "official total unsure"}:
            assert r["Note"], r["Paper ID"]


def test_q_numbers_reference_blocks():
    # q numbers must be parseable (numbers, optionally prefixed A/B)
    for r in _sheet("Questions"):
        assert re.fullmatch(r"[AB]?\d+|\?", str(r["Q no."])), r
