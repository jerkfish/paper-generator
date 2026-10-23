from __future__ import annotations

import re

import docx
from docx.oxml.ns import qn

from src.mat133.build_docx import build_paper_docx, build_scheme_docx, validate_docx
from src.mat133.models import Mat133Paper, Mat133Part, Mat133Question


def _table_text(tables) -> str:
    parts = []
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
                parts.append(_table_text(cell.tables))
    return "\n".join(parts)


def _paper() -> Mat133Paper:
    return Mat133Paper(
        month="NOVEMBER", year=2026,
        questions=[Mat133Question(number=1, parts=[
            Mat133Part(
                label="a", text="Solve x + 1 = 3.", marks=3,
                answer="x = 2", answer_latex="x = 2",
                solution_steps=["Subtract 1 from both sides."],
                solution_latex=["x + 1 = 3", "x = 3 - 1", "x = 2"],
                marking_notes=["M1 rearrange", "A1 answer"],
            ),
        ])],
    )


def test_question_headings_have_no_rule(tmp_path) -> None:
    paper = _paper()
    path = build_paper_docx(paper, tmp_path / "paper.docx")
    doc = docx.Document(str(path))
    headings = [p for p in doc.paragraphs
                if p.text.strip().startswith("QUESTION ")]
    assert headings
    for heading in headings:
        assert heading._p.find(".//" + qn("w:pBdr")) is None


def test_scheme_has_marker_cover_and_latex_images(tmp_path) -> None:
    paper = _paper()
    path = build_scheme_docx(paper, tmp_path / "scheme.docx")
    doc = docx.Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "INSTRUCTIONS TO MARKERS" in text
    assert "Do not bring any material" not in text
    assert "Answers and marking points are for local study use only." in text
    assert doc.tables
    table_text = "".join(re.findall(
        r"<w:t[^>]*>([^<]*)</w:t>", doc.element.body.xml))
    assert "SOLUTION" in table_text
    assert "Working:" not in table_text
    assert "Subtract 1 from both sides." not in table_text
    assert "The denominator cannot be zero" not in table_text
    assert "Now solve" not in table_text
    assert not any(line.strip().startswith(("1. ", "2. ", "3. "))
                   for line in table_text.splitlines())
    assert "M1" in table_text or "A1" in table_text
    # Crest plus answer/working formula images.
    media = [r for r in doc.part.rels.values()
             if "image" in r.target_ref]
    assert len(media) >= 4
    check = validate_docx(path, paper, scheme=True)
    assert check["xml_valid"]
    assert check["forbidden_tokens"] == []
    assert check["part_labels_found"] == 1
