"""Step 2 done-when tests:
every paper has a row in the Papers sheet, and uncertain fields are listed.
"""
from __future__ import annotations

import hashlib
import re

from openpyxl import load_workbook

from src.config import ORIGINAL_PAPERS_DIR, PAPERS_DIR, TRACKING_XLSX

ERA_VALUES = {
    "classic (provisional)", "transitional (provisional)",
    "modern (provisional)", "unsure",
}
EXTRACTED_FIELDS = {"University", "Course code", "Official total marks", "Duration", "Era"}


def _rows() -> list[dict]:
    assert TRACKING_XLSX.exists(), "run 'uv run python -m src.ingest.register_papers' first"
    ws = load_workbook(TRACKING_XLSX)["Papers"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def _md5(path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def _paper_files(directory) -> set[str]:
    return {p.name for p in directory.iterdir() if p.suffix.lower() == ".pdf"}


def test_every_paper_has_a_row():
    rows = _rows()
    names = _paper_files(PAPERS_DIR)
    assert len(rows) == len(names)
    assert {r["File name"] for r in rows} == names


def test_paper_ids_unique_and_wellformed():
    ids = [r["Paper ID"] for r in _rows()]
    assert len(ids) == len(set(ids))
    for pid in ids:
        assert re.fullmatch(r"CSC662-\d{4}-\d{2}(-copy)?", pid), pid


def test_month_and_year_match_file_name():
    for r in _rows():
        m = re.match(r"CSC662\s+([A-Z]+)\s+(\d{4})", r["File name"], re.I)
        assert m, r["File name"]
        assert r["Month"].upper() == m.group(1).upper()
        assert r["Year"] == int(m.group(2))


def test_extracted_fields_use_controlled_values():
    for r in _rows():
        assert r["University"] in {"Universiti Teknologi MARA", "unsure"}
        assert r["Course code"] in {"CSC662", "unsure"}
        assert r["Duration"] in {"3 hours", "unsure"}
        assert r["Era"] in ERA_VALUES
        total = r["Official total marks"]
        assert total == "unsure" or (isinstance(total, int) and 10 <= total <= 200)


def test_era_is_consistent_with_year():
    for r in _rows():
        era, year = r["Era"], r["Year"]
        if era.startswith("classic"):
            assert year <= 2012, r
        elif era.startswith("transitional"):
            assert 2013 <= year <= 2017, r
        elif era.startswith("modern"):
            assert year >= 2018, r


def test_unsure_fields_are_listed_in_notes():
    for r in _rows():
        if any(r[f] == "unsure" for f in EXTRACTED_FIELDS):
            assert "unsure" in str(r["Notes"]).lower(), r["File name"]


def test_copies_match_originals_byte_for_byte():
    for p in PAPERS_DIR.iterdir():
        if p.suffix.lower() != ".pdf":
            continue
        src = ORIGINAL_PAPERS_DIR / p.name
        assert src.exists(), f"original missing: {p.name}"
        assert _md5(src) == _md5(p), f"copy differs from original: {p.name}"


def test_every_original_is_copied():
    assert _paper_files(ORIGINAL_PAPERS_DIR) == _paper_files(PAPERS_DIR)


def test_duplicate_copies_are_flagged():
    for r in _rows():
        if "(2)" in r["File name"]:
            assert "byte-identical copy of" in str(r["Notes"]), r["File name"]
