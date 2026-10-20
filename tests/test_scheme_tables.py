"""Table geometry and PDF right-margin QA."""
from __future__ import annotations

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from src.assemble.table_fit import (body_usable_width_twips,
                                   fix_tables, grid_errors,
                                   pdf_right_margin_errors)


def _doc(tmp_path):
    d = Document()
    s = d.sections[0]
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(2.0)
    d.add_paragraph("test")
    t = d.add_table(rows=2, cols=6)
    for cell in t.rows[0].cells:
        cell.text = "header"
    d.save(str(tmp_path / "t.docx"))
    return d


def test_fix_tables_sets_exact_usable_width(tmp_path):
    d = _doc(tmp_path)
    usable = fix_tables(d)
    t = d.tables[0]
    grid = [int(c.get(qn("w:w")))
            for c in t._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(grid) == usable
    assert int(t._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) == usable
    assert t._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    assert t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(r._tr.trPr.find(qn("w:cantSplit")) is not None
               for r in t.rows)
    assert grid_errors(d) == []


def test_grid_overflow_is_detected(tmp_path):
    d = _doc(tmp_path)
    usable = fix_tables(d)
    # Force a grid sum above the usable width.
    cols = d.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))
    cols[0].set(qn("w:w"), str(int(cols[0].get(qn("w:w"))) + 500))
    errors = grid_errors(d, usable)
    assert any("gridCol sum" in e for e in errors)


def test_pdf_ink_right_of_margin_is_detected(tmp_path):
    pdf = tmp_path / "margin.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595.5, height=842.5)
    page.insert_text((540, 100), "overflow", fontsize=12)
    doc.save(str(pdf))
    errors = pdf_right_margin_errors(pdf, right_margin_twips=1320,
                                     tolerance_px=0)
    assert any("beyond right margin" in e for e in errors)
